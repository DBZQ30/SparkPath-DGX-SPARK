"""年级分区测试（v6）：db 层 + 主年级众数 + 上传校验。"""
import os
import time
import pytest
from fastapi.testclient import TestClient

from academicwarning.db import WarningDB
from academicwarning.models import SourceFile, Selection

# 上传校验测试鉴权（与 test_api.py 同模式）
TEST_KEY = "test-key-123"


def _db(tmp_path):
    db = WarningDB(str(tmp_path / "t.db"))
    db.init_schema()
    return db


def _sf(ftype, fname, grade):
    return SourceFile(file_type=ftype, file_name=fname, file_hash=fname,
                      file_path="/tmp/x", upload_time="2026-09-01 00:00:00",
                      uploader="t", grade=grade,
                      in_file_meta={"major": "工商管理"})


def test_init_schema_grade_columns(tmp_path):
    db = _db(tmp_path)
    cols = {r[1] for r in db.conn.execute("PRAGMA table_info(source_file)")}
    assert "grade" in cols
    cols = {r[1] for r in db.conn.execute("PRAGMA table_info(selection)")}
    assert "grade" in cols
    # init_schema 幂等：再跑一次不报错
    db.init_schema()


def test_grade_master_add_and_list(tmp_path):
    db = _db(tmp_path)
    db.add_grade("2023级")
    db.add_grade("2026级")
    db.add_grade("2023级")   # 重复忽略
    assert db.list_grades() == ["2023级", "2026级"]


def test_latest_source_file_by_grade(tmp_path):
    db = _db(tmp_path)
    db.insert_source_file(_sf("selection", "a.xlsx", "2023级"))
    db.insert_source_file(_sf("selection", "b.xlsx", "2024级"))
    assert db.latest_source_file("selection", "2024级").file_name == "b.xlsx"
    assert db.latest_source_file("selection", "2023级").file_name == "a.xlsx"
    assert db.latest_source_file("selection").file_name == "b.xlsx"  # 空=全取
    # F3：读回 grade 映射（_sf_from_row 显式列位）
    assert db.latest_source_file("selection", "2024级").grade == "2024级"
    assert db.latest_source_file("selection", "2023级").grade == "2023级"


def test_latest_active_plans_by_grade(tmp_path):
    db = _db(tmp_path)
    db.insert_plan_meta("工商管理", "v1", "p.docx", "2026-09-01 00:00:00",
                        grade="2023级")
    db.insert_plan_meta("工商管理", "v1", "p.docx", "2026-09-01 00:00:00",
                        grade="2026级")
    assert set(db.latest_active_plans("2023级")) == {"工商管理"}
    assert set(db.latest_active_plans("2026级")) == {"工商管理"}


def test_insert_get_selections_with_grade(tmp_path):
    """insert_selections/get_selections 带 grade 列（round-trip）。"""
    db = _db(tmp_path)
    db.insert_selections([
        Selection(student_id="S1", name="甲", major="工商管理", grade="2023级",
                  semester_label="2026-2027学年 第一学期", semester_code="4-1",
                  course_code="C1", course_name="高等数学", credit=4.0,
                  nature="必修", category="公共课程", status="选中",
                  source_file_id=7),
    ])
    got = db.get_selections(7)
    assert len(got) == 1
    assert got[0].grade == "2023级"
    assert got[0].name == "甲" and got[0].major == "工商管理" and got[0].credit == 4.0
    assert got[0].course_code == "C1" and got[0].status == "选中"
    # 无 grade 的旧行兼容（默认 ""）
    db.insert_selections([
        Selection(student_id="S2", name="乙", major="工业工程",
                  semester_label="2026-2027学年 第一学期", semester_code="4-1",
                  course_code="C2", course_name="英语", credit=2.0,
                  nature="必修", category="公共课程", status="选中",
                  source_file_id=8),
    ])
    assert db.get_selections(8)[0].grade == ""


def test_selection_main_grade_mode():
    """主年级 = 年级列众数；混级文件按众数。"""
    from academicwarning.service import selection_main_grade
    rows = [{"grade": "2023级"}] * 113 + [{"grade": "2022级"}] * 16
    assert selection_main_grade(rows) == "2023级"
    assert selection_main_grade([{"grade": ""}]) == ""


# ---- 上传年级校验（TestClient，模式同 test_api.py）----
from academicwarning import api


@pytest.fixture
def api_env(tmp_path, monkeypatch):
    """api 测试环境：临时库 + 上传目录 + 固定 key（同 test_api.py 的 autouse 模式）。"""
    monkeypatch.setattr(api, "_API_KEY", TEST_KEY)
    monkeypatch.setattr(api, "_UPLOAD_DIR", tmp_path / "uploads")
    db_path = str(tmp_path / "warning.db")
    import academicwarning.service as svc
    monkeypatch.setattr(api, "WarningDB", lambda: WarningDB(db_path))
    monkeypatch.setattr(svc, "WarningDB", lambda: WarningDB(db_path))
    _db = WarningDB(db_path)
    try:
        _db.init_schema()
    finally:
        _db.close()
    return db_path


@pytest.fixture
def client():
    return TestClient(api.app)


def _selection_xlsx(tmp_path, grade="2023级"):
    """最小选课 xlsx：表头含"学年学期"（detect_file_type 识别键）+ "年级"列。"""
    import openpyxl
    p = tmp_path / "选课结果.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["学年学期", "学号", "姓名", "年级", "专业名称", "课程号", "课程名",
               "课程性质", "课程类别", "选课状态", "重修重考"])
    ws.append(["2026-2027学年 第一学期", "2023001", "张三", grade, "工商管理",
               "C1", "高等数学", "必修", "公共课程", "选中", "初修"])
    wb.save(p)
    return p


def _wait_done(db_path, file_type, grade, timeout=15.0):
    """轮询 db 直到指定年级分区出现 done 记录（后台线程解析）。"""
    deadline = time.time() + timeout
    while time.time() < deadline:
        db = WarningDB(db_path)
        try:
            sf = db.latest_source_file(file_type, grade)
        finally:
            db.close()
        if sf and sf.parsed_status == "done":
            return sf
        time.sleep(0.2)
    raise AssertionError(f"{file_type}({grade}) 解析超时")


def test_upload_selection_wrong_grade_rejected(client, api_env, tmp_path):
    """v6 防传错：选课文件主年级 2023级，选 2024级 入口 → 同步拒绝。"""
    f = _selection_xlsx(tmp_path, grade="2023级")
    with open(f, "rb") as fh:
        r = client.post(
            "/api/warning/upload",
            files={"file": ("选课结果.xlsx", fh, "application/octet-stream")},
            data={"type_hint": "选课结果", "grade_hint": "2024级"},
            headers={"X-API-Key": TEST_KEY},
        )
    assert r.status_code == 200
    body = r.json()
    assert body["parsed_status"] == "rejected"
    assert "2023级" in body["message"] and "不匹配" in body["message"]
    # 库中无残留记录
    db = WarningDB(api_env)
    assert db.latest_source_file("selection", "2024级") is None
    db.close()


def test_upload_selection_matching_grade_parsed(client, api_env, tmp_path):
    """v6：主年级与所选一致 → 接收；后台解析后 selection/source_file 落 grade。"""
    f = _selection_xlsx(tmp_path, grade="2023级")
    with open(f, "rb") as fh:
        r = client.post(
            "/api/warning/upload",
            files={"file": ("选课结果.xlsx", fh, "application/octet-stream")},
            data={"type_hint": "选课结果", "grade_hint": "2023级"},
            headers={"X-API-Key": TEST_KEY},
        )
    body = r.json()
    assert body["parsed_status"] == "parsing"
    assert body["grade"] == "2023级"
    sf = _wait_done(api_env, "selection", "2023级")
    assert sf.grade == "2023级"
    db = WarningDB(api_env)
    try:
        got = db.get_selections(sf.id)
    finally:
        db.close()
    assert got and got[0].grade == "2023级"
    assert got[0].student_id == "2023001"


@pytest.mark.samples
def test_upload_plan_filename_grade_mismatch_rejected(client, api_env, tmp_path):
    """v6 防传错：plan 文件名年级 2024级，选 2023级 入口 → 同步拒绝。"""
    src = "academicwarning/docs/2023版工商管理专业培养方案.docx"
    if not os.path.exists(src):
        pytest.skip("样本缺失")
    with open(src, "rb") as f:
        r = client.post(
            "/api/warning/upload",
            files={"file": ("2024级工商管理培养方案.docx", f,
                            "application/octet-stream")},
            data={"type_hint": "培养方案", "major_hint": "工商管理",
                  "grade_hint": "2023级"},
            headers={"X-API-Key": TEST_KEY},
        )
    body = r.json()
    assert body["parsed_status"] == "rejected"
    assert "2024级" in body["message"] and "不匹配" in body["message"]


@pytest.mark.samples
def test_upload_plan_matching_grade_parsed(client, api_env, tmp_path):
    """v6：plan 文件名年级与所选一致 → 接收；training_plan 落 grade。"""
    src = "academicwarning/docs/2023版工商管理专业培养方案.docx"
    if not os.path.exists(src):
        pytest.skip("样本缺失")
    with open(src, "rb") as f:
        r = client.post(
            "/api/warning/upload",
            files={"file": ("2023级工商管理培养方案.docx", f,
                            "application/octet-stream")},
            data={"type_hint": "培养方案", "major_hint": "工商管理",
                  "grade_hint": "2023级"},
            headers={"X-API-Key": TEST_KEY},
        )
    body = r.json()
    assert body["parsed_status"] == "parsing"
    assert body["grade"] == "2023级"
    _wait_done(api_env, "plan", "2023级")
    db = WarningDB(api_env)
    try:
        row = db.conn.execute(
            "SELECT grade FROM training_plan WHERE major_name='工商管理'"
            " ORDER BY id DESC LIMIT 1").fetchone()
    finally:
        db.close()
    assert row and row[0] == "2023级"


def test_upload_invalid_grade_format_rejected(client, api_env, tmp_path):
    """v6：grade_hint 非 20xx级 格式 → 拒绝。"""
    f = _selection_xlsx(tmp_path, grade="2023级")
    with open(f, "rb") as fh:
        r = client.post(
            "/api/warning/upload",
            files={"file": ("选课结果.xlsx", fh, "application/octet-stream")},
            data={"type_hint": "选课结果", "grade_hint": "23级"},
            headers={"X-API-Key": TEST_KEY},
        )
    body = r.json()
    assert body["parsed_status"] == "rejected"
    assert "年级格式无效" in body["message"]


# ---- Task 3：年级状态 / 删除分区 / 添加年级端点 ----

def test_add_grade_endpoint_validation(client, api_env):
    """Task 3：POST /api/warning/grade 格式校验 + 重复拒绝（ok=false+message）。"""
    r = client.post("/api/warning/grade", json={"name": "2026级"},
                    headers={"X-API-Key": TEST_KEY})
    assert r.status_code == 200
    assert r.json()["ok"] is True
    # 格式错（非 20xx级）
    r2 = client.post("/api/warning/grade", json={"name": "26级"},
                     headers={"X-API-Key": TEST_KEY})
    assert r2.json()["ok"] is False
    assert "格式" in r2.json()["message"]
    # 重复添加拒绝（INSERT OR IGNORE 静默，需 api 层先查）
    r3 = client.post("/api/warning/grade", json={"name": "2026级"},
                     headers={"X-API-Key": TEST_KEY})
    assert r3.json()["ok"] is False
    assert "已存在" in r3.json()["message"]
    db = WarningDB(api_env)
    try:
        assert db.list_grades() == ["2026级"]
    finally:
        db.close()


def test_add_grade_requires_auth(client):
    """Task 3：添加年级需鉴权（同其他端点，401）。"""
    r = client.post("/api/warning/grade", json={"name": "2026级"})
    assert r.status_code == 401


def test_status_returns_grades(client, api_env):
    """Task 3：/status 响应顶层带 grades（前端 tab 数据源）。"""
    db = WarningDB(api_env)
    try:
        db.add_grade("2023级")
        db.add_grade("2026级")
    finally:
        db.close()
    r = client.get("/api/warning/status", headers={"X-API-Key": TEST_KEY})
    assert r.status_code == 200
    assert r.json()["grades"] == ["2023级", "2026级"]


def test_delete_by_grade_partition(client, api_env):
    """Task 3：删除带 grade → 仅删该年级分区，其他年级保留。"""
    db = WarningDB(api_env)
    try:
        db.insert_source_file(_sf("selection", "a.xlsx", "2023级"))
        db.insert_source_file(_sf("selection", "b.xlsx", "2024级"))
    finally:
        db.close()
    r = client.post("/api/warning/delete",
                    json={"file_type": "selection", "grade": "2023级"},
                    headers={"X-API-Key": TEST_KEY})
    assert r.status_code == 200
    assert "已删除" in r.json()["message"]
    db = WarningDB(api_env)
    try:
        assert db.latest_source_file("selection", "2023级") is None
        assert db.latest_source_file("selection", "2024级").file_name == "b.xlsx"
    finally:
        db.close()


def test_delete_plan_by_grade_keeps_other_grade(client, api_env):
    """Task 3：plan 按 年级×专业 删除 → 级联只删该年级 training_plan，
    其他年级的方案/文件保留（分区隔离）。"""
    db = WarningDB(api_env)
    try:
        db.insert_plan_meta("工商管理", "v1", "p2023.docx", "2026-09-01 00:00:00",
                            grade="2023级")
        db.insert_plan_meta("工商管理", "v1", "p2026.docx", "2026-09-01 00:00:00",
                            grade="2026级")
        db.insert_source_file(_sf("plan", "p2023.docx", "2023级"))
        db.insert_source_file(_sf("plan", "p2026.docx", "2026级"))
    finally:
        db.close()
    r = client.post("/api/warning/delete",
                    json={"file_type": "plan", "major": "工商管理",
                          "grade": "2023级"},
                    headers={"X-API-Key": TEST_KEY})
    assert r.status_code == 200
    assert "已删除" in r.json()["message"]
    db = WarningDB(api_env)
    try:
        assert db.latest_source_file("plan", "2023级") is None
        assert db.latest_source_file("plan", "2026级") is not None
        assert "工商管理" not in db.latest_active_plans("2023级")
        assert "工商管理" in db.latest_active_plans("2026级")
    finally:
        db.close()


# ---- Task 4：按年级独立检查 ----

def _seed_check_by_grade(db):
    """两份年级数据（2023级 S1 触发 / 2024级 S2 修满）——service/API 测试共用。"""
    from academicwarning.models import PlanCourse, Selection, Grade
    from academicwarning.parsers import clean_course_name

    def _plan(grade: str, semester: str) -> None:
        db.insert_plan_meta("工商管理", "v1", f"p-{grade}.docx",
                            "2026-09-01 00:00:00", grade=grade, elective_req=8.0)
        pid = db.conn.execute(
            "SELECT id FROM training_plan WHERE major_name='工商管理'"
            " AND grade=? ORDER BY id DESC LIMIT 1", (grade,)).fetchone()[0]
        db.insert_plan_courses(pid, [
            PlanCourse(plan_id=pid, course_code=f"EL{i}", course_name=f"选修课{i}",
                       credit=2.0, course_type="专业选修课程", required_flag="选修",
                       semester=semester)
            for i in range(1, 5)])

    def _sel_file(grade: str, entry_year: int, semester_code: str,
                  sid: str, name: str) -> None:
        fid = db.insert_source_file(SourceFile(
            file_type="selection", file_name=f"s-{grade}.xlsx", file_hash=f"h-{grade}",
            file_path="", upload_time="2026-09-01 00:00:00", uploader="t",
            parsed_status="done", grade=grade,
            in_file_meta={"semester_label": "2026-2027学年 第一学期",
                          "semester_code": semester_code,
                          "grades": [grade], "entry_year": entry_year}))
        db.insert_selections([Selection(
            student_id=sid, name=name, major="工商管理", grade=grade,
            semester_label="2026-2027学年 第一学期", semester_code=semester_code,
            course_code="C1", course_name="公共课", credit=2.0, nature="必修",
            category="公共课程", status="选中", source_file_id=fid)])

    def _grade_file(grade: str, rows: list) -> None:
        gid = db.insert_source_file(SourceFile(
            file_type="grade", file_name=f"g-{grade}.docx", file_hash=f"hg-{grade}",
            file_path="", upload_time="2026-09-01 00:00:00", uploader="t",
            parsed_status="done", grade=grade, in_file_meta={"major": "工商管理"}))
        for g in rows:
            g.source_file_id = gid
        db.insert_grades(rows)

    def _roster_file(grade: str, sid: str, name: str) -> None:
        # v6.5：名单 = 学籍 roster（在籍在校过滤已由 parse_roster 完成）
        rid = db.insert_source_file(SourceFile(
            file_type="roster", file_name=f"r-{grade}.xls", file_hash=f"hr-{grade}",
            file_path="", upload_time="2026-09-01 00:00:00", uploader="t",
            parsed_status="done", grade=grade))
        db.insert_roster([{"student_id": sid, "name": name, "grade": grade,
                           "major": "工商管理", "class_name": "工商2301",
                           "status": "正常", "source_file_id": rid}])

    # 2023级：S1 无专业选修（应修 8，已修 0）→ 触发
    _plan("2023级", "3-1")
    _sel_file("2023级", 2023, "4-1", "S1", "甲")
    _grade_file("2023级", [
        Grade(student_id="S1", student_name="甲", course_name=f"通识课{i}◆",
              credit=2.0, grade_raw="80", pass_flag=1, marker="◆选修课",
              course_name_clean=clean_course_name(f"通识课{i}◆"))
        for i in range(6)])
    _roster_file("2023级", "S1", "甲")
    # 2024级：S2 修满（模块 12 + 专业选修 8）→ 不触发
    _plan("2024级", "3-1")
    _sel_file("2024级", 2024, "3-1", "S2", "乙")
    _grade_file("2024级", [
        Grade(student_id="S2", student_name="乙", course_name=f"通识课{i}◆",
              credit=2.0, grade_raw="80", pass_flag=1, marker="◆选修课",
              course_name_clean=clean_course_name(f"通识课{i}◆"))
        for i in range(6)] + [
        Grade(student_id="S2", student_name="乙", course_name=f"选修课{i}",
              credit=2.0, grade_raw="85", pass_flag=1,
              course_name_clean=clean_course_name(f"选修课{i}"))
        for i in range(1, 5)])
    _roster_file("2024级", "S2", "乙")


def test_run_selection_check_by_grade(tmp_path):
    """2023级 与 2024级 各自独立检查，互不串数据。

    2023级：S1 选课（无专业选修）→ 触发；2024级：S2 选课（修满）→ 不触发；
    断言 selection_check 落库行的 grade。"""
    from academicwarning.service import run_selection_check
    db = _db(tmp_path)
    _seed_check_by_grade(db)

    n1, _ = run_selection_check(db, "2023级")
    assert n1 == 1   # S1 触发
    n2, _ = run_selection_check(db, "2024级")
    assert n2 == 0   # S2 修满不触发
    rows = db.conn.execute(
        "SELECT student_id, grade FROM selection_check ORDER BY grade").fetchall()
    assert rows == [("S1", "2023级")]   # 落库行带 grade、按年级隔离


def test_selection_check_api_by_grade(client, api_env):
    """Task 4：selection-check 接口带年级——POST run / GET / status 按年级隔离。"""
    db = WarningDB(api_env)
    try:
        _seed_check_by_grade(db)
    finally:
        db.close()

    r = client.post("/api/warning/selection-check/run", json={"grade": "2023级"},
                    headers={"X-API-Key": TEST_KEY})
    assert r.json()["count"] == 1
    r = client.post("/api/warning/selection-check/run", json={"grade": "2024级"},
                    headers={"X-API-Key": TEST_KEY})
    assert r.json()["count"] == 0

    r = client.get("/api/warning/selection-check?grade=2023级",
                   headers={"X-API-Key": TEST_KEY})
    assert r.json()["grade"] == "2023级"
    assert [s["student_id"] for s in r.json()["students"]] == ["S1"]
    r = client.get("/api/warning/selection-check?grade=2024级",
                   headers={"X-API-Key": TEST_KEY})
    assert r.json()["grade"] == "2024级"
    assert r.json()["students"] == []

    # 不带 grade（旧客户端）→ 最新文件（2024级）
    r = client.get("/api/warning/selection-check", headers={"X-API-Key": TEST_KEY})
    assert r.json()["grade"] == ""
    assert r.json()["source_file_id"] is not None

    # /status 按年级隔离 selection 行
    def _sel_row(grade):
        r = client.get(f"/api/warning/status?grade={grade}",
                       headers={"X-API-Key": TEST_KEY})
        return next(f for f in r.json()["files"]
                    if f["file_type"] == "selection")["items"][0]
    assert _sel_row("2023级")["file_name"] == "s-2023级.xlsx"
    assert _sel_row("2024级")["file_name"] == "s-2024级.xlsx"


# ---- Task 5：通识课选课学分补全（孙镜麒误报根因）----

def _grade(sid, name, credit, pass_flag=1, marker=""):
    from academicwarning.models import Grade
    from academicwarning.parsers import clean_course_name
    return Grade(student_id=sid, course_name=name, credit=credit, grade_raw="80",
                 pass_flag=pass_flag, marker=marker,
                 course_name_clean=clean_course_name(name))


def _sel(sid, code, name, category="专业选修课程", credit=2.0):
    return Selection(student_id=sid, semester_label="2026-2027学年 第一学期",
                     semester_code="4-1", course_code=code, course_name=name,
                     credit=credit, nature="选修", category=category,
                     status="选中", retake="初修")


def test_gen_ed_credit_override():
    """通识选课：本人成绩单优先，其次全校众数，最后 2.0 估算。"""
    from academicwarning.service import gen_ed_credit_override
    grades = [_grade("S1", "道家的智慧", 2.0), _grade("S2", "核与世界", 2.0)]
    sels = [_sel("S1", "CORE100210", "道家的智慧", category="基础通识类核心课", credit=0.0),
            _sel("S2", "CORE100210", "道家的智慧", category="基础通识类核心课", credit=0.0),
            _sel("S1", "GNED100001", "人工智能与新媒体文化", category="基础通识类选修课", credit=0.0)]
    gen_ed_credit_override(grades, sels)
    assert sels[0].credit == 2.0   # 本人成绩单
    assert sels[1].credit == 2.0   # 全校众数（S1 的 2.0）
    assert sels[2].credit == 2.0   # 无记录 → 估算
    # 优先级区分：S2 有本人同名 3.0 记录（全校众数 2.0）→ 取本人 3.0
    sels3 = [_sel("S2", "CORE100210", "道家的智慧", category="基础通识类核心课", credit=0.0)]
    gen_ed_credit_override([_grade("S1", "道家的智慧", 2.0),
                            _grade("S2", "道家的智慧", 3.0)], sels3)
    assert sels3[0].credit == 3.0
    # 非通识类 / 已有学分的选课不受影响
    sels4 = [_sel("S1", "EL1", "选修课A", category="专业选修课程", credit=2.0),
             _sel("S1", "GNED1", "人工智能与新媒体文化", category="基础通识类选修课", credit=1.5)]
    gen_ed_credit_override(grades, sels4)
    assert sels4[0].credit == 2.0 and sels4[1].credit == 1.5
    # 清洗后为空名的选课行跳过（F3：避免误匹配无关成绩行）
    sels5 = [_sel("S1", "GNED0", "◆", category="基础通识类选修课", credit=0.0)]
    gen_ed_credit_override(grades, sels5)
    assert sels5[0].credit == 0.0


# ---- 2026-09-07 D5：通识课程信息表作学分权威（service._GEN_ED_TABLE 表驱动）----

def _write_gen_ed_table(path, rows):
    """tmp 造最小通识表（表头与 docs/通识课程信息表.xlsx 一致；rows=(课号, 课名, 学分)）。"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "基本信息"
    ws.append(["序号", "课程号", "课程名", "课程模块", "开课单位", "类型",
               "课程负责人", "备注", "学分"])
    for i, (code, name, credit) in enumerate(rows, 1):
        ws.append([i, code, name, "科学探索与技术创新", "某学院", "核心课",
                   "张三", None, credit])
    wb.save(path)


def test_gen_ed_table_code_then_name_then_mode(tmp_path, monkeypatch):
    """D5 表驱动顺序：本人无同名时，表课程号 → 表 clean 课名 → 全校众数 → 2.0。"""
    import academicwarning.service as svc
    from academicwarning.service import gen_ed_credit_override
    tbl = tmp_path / "通识课程信息表.xlsx"
    # 同名两行：GNED700001 先入 by_name（首行优先 5.0）；CORE100313 由课程号命中 2.0
    _write_gen_ed_table(tbl, [
        ("GNED700001", "纳米科技前沿", 5.0),
        ("CORE100313", "纳米科技前沿", 2.0),
    ])
    monkeypatch.setattr(svc, "_GEN_ED_TABLE_FILE", tbl)
    monkeypatch.setattr(svc, "_GEN_ED_TABLE", None)
    grades = [_grade("S9", "纳米科技前沿", 3.0), _grade("S9", "核与世界", 3.0)]
    sels = [_sel("S1", "CORE100313", "纳米科技前沿", category="基础通识类核心课", credit=0.0),
            _sel("S2", "GNED999901", "纳米科技前沿", category="基础通识类核心课", credit=0.0),
            _sel("S3", "GNED999902", "核与世界", category="基础通识类核心课", credit=0.0),
            _sel("S4", "GNED999903", "未见表新课", category="基础通识类核心课", credit=0.0)]
    gen_ed_credit_override(grades, sels)
    assert sels[0].credit == 2.0   # 表课程号（先于 by_name 5.0 与众数 3.0）
    assert sels[1].credit == 5.0   # 课号不在表 → 表 clean 课名（先于众数 3.0）
    assert sels[2].credit == 3.0   # 表外通识课 → 全校众数（原兜底链保留）
    assert sels[3].credit == 2.0   # 无任何命中 → 2.0 估算


def test_gen_ed_table_own_grade_first_clean_name(tmp_path, monkeypatch):
    """D5：本人成绩同名仍最优先（先于表课程号）；表内课名带前后空格经 clean 命中。"""
    import academicwarning.service as svc
    from academicwarning.service import gen_ed_credit_override
    tbl = tmp_path / "通识课程信息表.xlsx"
    _write_gen_ed_table(tbl, [("CORE100313", " 纳米科技前沿 ", 2.0)])
    monkeypatch.setattr(svc, "_GEN_ED_TABLE_FILE", tbl)
    monkeypatch.setattr(svc, "_GEN_ED_TABLE", None)
    grades = [_grade("S1", "纳米科技前沿", 3.0)]   # S1 本人同名 3.0（表内 2.0）
    sels = [_sel("S1", "CORE100313", "纳米科技前沿◆",
                 category="基础通识类核心课", credit=0.0)]
    gen_ed_credit_override(grades, sels)
    assert sels[0].credit == 3.0   # 本人成绩同名先于表课程号


def test_gen_ed_table_candidacy_expand_excluded_kept(tmp_path, monkeypatch):
    """D5：表命中扩展候选（课号在表即补全，类别非通识也可）；跨选/辅修仍排除。"""
    import academicwarning.service as svc
    from academicwarning.service import gen_ed_credit_override
    tbl = tmp_path / "通识课程信息表.xlsx"
    _write_gen_ed_table(tbl, [("CORE100313", "纳米科技前沿", 2.0)])
    monkeypatch.setattr(svc, "_GEN_ED_TABLE_FILE", tbl)
    monkeypatch.setattr(svc, "_GEN_ED_TABLE", None)
    sels = [_sel("S1", "CORE100313", "纳米科技前沿", category="专业选修课程", credit=0.0),
            _sel("S2", "CORE100313", "纳米科技前沿", category="跨选课", credit=0.0),
            _sel("S3", "CORE100313", "纳米科技前沿", category="辅修课程", credit=0.0)]
    gen_ed_credit_override([], sels)
    assert sels[0].credit == 2.0   # 课程号在表 → 非通识类别也补全
    assert sels[1].credit == 0.0   # 跨选课不参与（EXCLUDED_CATEGORIES）
    assert sels[2].credit == 0.0   # 辅修课程不参与


def test_gen_ed_table_missing_falls_back_no_raise(tmp_path, monkeypatch):
    """D5：通识表缺失 → 不抛异常阻断 run，回落众数/2.0 兜底链。"""
    import academicwarning.service as svc
    from academicwarning.service import gen_ed_credit_override
    monkeypatch.setattr(svc, "_GEN_ED_TABLE_FILE", tmp_path / "不存在.xlsx")
    monkeypatch.setattr(svc, "_GEN_ED_TABLE", None)
    grades = [_grade("S9", "道家的智慧", 3.0)]
    sels = [_sel("S1", "CORE100210", "道家的智慧",
                 category="基础通识类核心课", credit=0.0)]
    gen_ed_credit_override(grades, sels)   # 不应抛异常
    assert sels[0].credit == 3.0   # 众数兜底仍生效


def test_assemble_check_prep_persists_gen_ed_credits(tmp_path):
    """Task 5 接入：_assemble_check_prep 通识选课 credit 补全并落库（展示一致）。"""
    from academicwarning.models import PlanCourse, Grade
    from academicwarning.parsers import clean_course_name
    from academicwarning.service import _assemble_check_prep
    db = _db(tmp_path)
    db.insert_plan_meta("工商管理", "v1", "p.docx", "2026-09-01 00:00:00",
                        grade="2023级", elective_req=8.0)
    pid = db.conn.execute("SELECT id FROM training_plan WHERE major_name='工商管理'"
                          " AND grade='2023级' ORDER BY id DESC LIMIT 1").fetchone()[0]
    db.insert_plan_courses(pid, [PlanCourse(plan_id=pid, course_code="C1", course_name="高等数学",
                                            credit=4.0, course_type="公共课程",
                                            required_flag="必修", semester="1-1")])
    fid = db.insert_source_file(SourceFile(
        file_type="selection", file_name="s.xlsx", file_hash="h1",
        file_path="", upload_time="2026-09-01 00:00:00", uploader="t",
        parsed_status="done", grade="2023级",
        in_file_meta={"semester_label": "2026-2027学年 第一学期",
                      "semester_code": "4-1", "entry_year": 2023}))
    db.insert_selections([
        Selection(student_id="S1", name="甲", major="工商管理", grade="2023级",
                  semester_label="2026-2027学年 第一学期", semester_code="4-1",
                  course_code="CORE100210", course_name="道家的智慧", credit=0.0,
                  nature="选修", category="基础通识类核心课", status="选中",
                  source_file_id=fid),
    ])
    gid = db.insert_source_file(SourceFile(
        file_type="grade", file_name="g.docx", file_hash="hg1",
        file_path="", upload_time="2026-09-01 00:00:00", uploader="t",
        parsed_status="done", grade="2023级", in_file_meta={"major": "工商管理"}))
    db.insert_grades([Grade(student_id="S1", student_name="甲",
                            course_name="道家的智慧", credit=2.0, grade_raw="80",
                            pass_flag=1, course_name_clean=clean_course_name("道家的智慧"),
                            source_file_id=gid)])
    # v6.5：名单 = 学籍 roster
    rid = db.insert_source_file(SourceFile(
        file_type="roster", file_name="r.xls", file_hash="hr1",
        file_path="", upload_time="2026-09-01 00:00:00", uploader="t",
        parsed_status="done", grade="2023级"))
    db.insert_roster([{"student_id": "S1", "name": "甲", "grade": "2023级",
                       "major": "工商管理", "class_name": "工商2301",
                       "status": "正常", "source_file_id": rid}])
    sel_file = db.latest_source_file("selection", "2023级")
    prep = _assemble_check_prep(db, sel_file)
    assert prep is not None
    s = prep.selections[0]
    assert s.credit == 2.0   # 本人成绩单同名 → 补全
    assert s.id is not None  # get_selections 带 id（落库依据）
    row = db.conn.execute("SELECT credit FROM selection WHERE id=?", (s.id,)).fetchone()
    assert row and row[0] == 2.0   # 学分变化已落库


def test_gen_ed_selected_counts_toward_module_gap():
    """Task 5 回归（孙镜麒误报根因）：通识选课（学分已补全）计入模块课程已选——
    CATEGORY_MAP 值"模块课程·核心/选修课程"须归并到 JSON 类键"模块课程"。"""
    from academicwarning.models import Student, PlanCourse
    from academicwarning.rules import Prep
    from academicwarning.selection_check import check_selection_rationality
    students = [Student(student_id="S1", name="甲", major="0824工商管理")]
    courses = [PlanCourse(plan_id=1, course_code=f"EL{i}", course_name=f"选修课{n}",
                          credit=2.0, course_type="专业选修课程", required_flag="选修",
                          semester=sem)
               for i, (n, sem) in enumerate(
                   (("A", "3-1"), ("B", "3-2"), ("C", "3-1"), ("D", "3-2")), 1)]
    sels = [_sel("S1", "CORE100210", "道家的智慧", category="基础通识类核心课", credit=2.0),
            _sel("S1", "GNED100001", "人工智能与新媒体文化", category="基础通识类选修课", credit=2.0)]
    # 2026-09-07 D3 及格制后：已修按同课最高分代表行计——四门各异课各一行
    # （旧同名行 ×4 累加会被代表行折叠为 2.0，不再是"已修满 8"的合法构造）
    grades = ([_grade("S1", f"选修课{n}", 2.0) for n in "ABCD"]   # 专业选修已修满 8
              + [_grade("S1", f"通识课{i}◆", 2.0, marker="◆选修课")
                 for i in range(5)])              # 模块课程已修 10
    credit_req = {"工商管理": {"模块课程": 12.0, "专业选修课程": 8.0}}
    prep = Prep(students, sels, grades, {"工商管理": courses}, {}, [], "4-1",
                entry_year=2023, credit_req_by_major=credit_req)
    assert prep.selected_credits_by_category("S1")["模块课程"] == 4.0
    assert check_selection_rationality(prep) == []   # 12-10-4 < 0 → 不再误报


def test_math_science_selected_key_unified():
    """审查 F1：数学和基础科学类课程选课键归并到学科门类基础课程
    （与 _CAT_RENAME 同语义——应修侧已归并，已选侧此前漏归并）。"""
    from academicwarning.models import Student
    from academicwarning.rules import Prep
    prep = Prep([Student(student_id="S1", name="甲", major="0824工商管理")],
                [_sel("S1", "MATH298207", "高等数学", category="数学和基础科学类课程",
                      credit=6.5)],
                [], {}, {}, [], "4-1")
    got = prep.selected_credits_by_category("S1")
    assert got.get("学科门类基础课程", 0.0) == 6.5
    assert "数学和基础科学类课程" not in got


# ---- Task 6：全员无成绩豁免子串匹配（方案名⊂成绩单名，单向防前缀碰撞）
# ----       + 豁免清单前端通知 ----

def test_exempt_match_substring():
    """豁免判定子串单向（方案名 ⊂ 成绩单名）：成绩单《管理学新生导论课》命中
    方案《管理学新生导论》→ 不豁免（此前精确匹配误判全员无成绩 → 应修低估 +
    缺修必修不查，32 人误判）。"""
    from academicwarning.models import PlanCourse, Student
    from academicwarning.rules import Prep
    courses = [PlanCourse(plan_id=1, course_code="C1", course_name="管理学新生导论",
                          credit=1.0, course_type="专业大类基础课程",
                          required_flag="必修", semester="1-1")]
    grades = [_grade("S1", "管理学新生导论课", 1.0),
              _grade("S2", "管理学新生导论课", 1.0)]
    students = [Student(student_id="S1", name="甲", grade="2023级", major="工商管理"),
                Student(student_id="S2", name="乙", grade="2023级", major="工商管理")]
    prep = Prep(students, [], grades, {"工商管理": courses}, {}, [], "4-1",
                entry_year=2023)
    assert "C1" not in prep.missing_courses("S1")   # 有成绩 → 不豁免


def test_exempt_no_prefix_collision():
    """防前缀碰撞：成绩单《国际商务》不得让《国际商务研究前沿》免豁免——
    短名是长名前缀 ≠ 同一课程（双向匹配 64 假阳性的根因）；豁免判定只认
    "成绩单名包含方案名"（方案名 ⊂ 成绩单名）方向。"""
    from academicwarning.models import PlanCourse, Student
    from academicwarning.rules import Prep
    courses = [
        PlanCourse(plan_id=1, course_code="C1", course_name="国际商务",
                   credit=2.0, course_type="专业大类基础课程",
                   required_flag="必修", semester="2-1"),
        PlanCourse(plan_id=1, course_code="C2", course_name="国际商务研究前沿",
                   credit=2.0, course_type="专业大类基础课程",
                   required_flag="必修", semester="3-1"),
    ]
    grades = [_grade("S1", "国际商务", 2.0), _grade("S2", "国际商务", 2.0)]
    students = [Student(student_id="S1", name="甲", grade="2023级", major="工商管理"),
                Student(student_id="S2", name="乙", grade="2023级", major="工商管理")]
    prep = Prep(students, [], grades, {"工商管理": courses}, {}, [], "4-1",
                entry_year=2023)
    assert "C1" not in prep.missing_courses("S1")   # 国际商务有成绩 → 不豁免
    assert "C2" in prep.missing_courses("S1")        # 研究前沿无成绩 → 仍豁免


def test_match_longest_clean_name():
    """v6.1 回归：方案含《管理学》与《管理学新生导论》，成绩《管理学新生导论课》
    应归因到 clean 名最长者《管理学新生导论》（此前按 dict 顺序命中"管理学" →
    工业工程缺修判定全员假阳性，被豁免修复暴露）。"""
    from academicwarning.models import PlanCourse, Student
    from academicwarning.rules import Prep
    from academicwarning.selection_check import _missing_required_courses
    courses = [
        PlanCourse(plan_id=1, course_code="MAGT478108", course_name="管理学",
                   credit=2.0, course_type="专业大类基础课程",
                   required_flag="必修", semester="1-2"),
        PlanCourse(plan_id=1, course_code="MAGT072408", course_name="管理学新生导论",
                   credit=1.0, course_type="专业大类基础课程",
                   required_flag="必修", semester="1-1"),
    ]
    students = [Student(student_id="S1", name="甲", grade="2023级", major="工商管理"),
                Student(student_id="S2", name="乙", grade="2023级", major="工商管理")]
    grades = [_grade("S1", "管理学新生导论课", 1.0),
              _grade("S2", "管理学新生导论课", 1.0)]
    prep = Prep(students, [], grades, {"工商管理": courses}, {}, [], "4-1",
                entry_year=2023)
    m = prep._match_plan_course("S1", grades[0])
    assert m is not None and m.course_code == "MAGT072408"   # 最长 clean 名胜出
    missing_req = [n for _, n, _, _, _, _ in _missing_required_courses(prep, "S1")]
    assert missing_req == ["管理学"]                          # 缺修判定不含《管理学新生导论》


def test_exempt_courses_list():
    """豁免清单：{专业: [课程名, ...]}（前端显著通知管理员核对），空时返回空 dict。"""
    from academicwarning.models import PlanCourse, Student
    from academicwarning.rules import Prep
    from academicwarning.selection_check import exempt_courses
    courses = [PlanCourse(plan_id=1, course_code="C1", course_name="管理学新生导论",
                          credit=1.0, course_type="专业大类基础课程",
                          required_flag="必修", semester="1-1")]
    students = [Student(student_id="S1", name="甲", grade="2023级", major="工商管理"),
                Student(student_id="S2", name="乙", grade="2023级", major="工商管理")]
    # 专业有成绩记录（covered 非空）但课程无人修 → 清单含该课程名
    prep = Prep(students, [], [_grade("S1", "高等数学", 4.0)],
                {"工商管理": courses}, {}, [], "4-1", entry_year=2023)
    assert exempt_courses(prep) == {"工商管理": ["管理学新生导论"]}
    # 课程有成绩 → 空 dict
    prep2 = Prep(students, [], [_grade("S1", "管理学新生导论", 1.0)],
                 {"工商管理": courses}, {}, [], "4-1", entry_year=2023)
    assert exempt_courses(prep2) == {}
    # 2026-09-01 用户反馈：专业选修/模块课程（JSON 全量口径）的豁免是噪音 → 不进清单
    elect = [PlanCourse(plan_id=1, course_code="E1", course_name="深度学习",
                        credit=2.0, course_type="专业选修课程",
                        required_flag="选修", semester="4-1")]
    prep3 = Prep(students, [], [_grade("S1", "高等数学", 4.0)],
                 {"工商管理": elect}, {}, [], "4-1", entry_year=2023)
    assert exempt_courses(prep3) == {}


# ---- Task 7：检查后自动导出 xlsx + 下载端点 ----

def test_build_report_xlsx(tmp_path, monkeypatch):
    """导出文件存在且含汇总表头；文件名带年级。"""
    from academicwarning.models import Student, PlanCourse
    from academicwarning.rules import Prep
    from academicwarning.selection_check import check_selection_rationality
    import academicwarning.export as exp
    monkeypatch.setattr(exp, "OUT_DIR", str(tmp_path))   # 落盘到临时目录，不污染仓库
    courses = [PlanCourse(plan_id=1, course_code="EL1", course_name="选修课A",
                          credit=2.0, course_type="专业选修课程", required_flag="选修",
                          semester="3-1")]
    students = [Student(student_id="S1", name="甲", grade="2023级", major="工商管理")]
    sels = [_sel("S1", "EL1", "选修课A", credit=2.0)]
    grades = [_grade("S1", "通识课1◆", 2.0, marker="◆选修课")]
    credit_req = {"工商管理": {"专业选修课程": 8.0}}
    prep = Prep(students, sels, grades, {"工商管理": courses}, {}, [], "4-1",
                entry_year=2023, credit_req_by_major=credit_req)
    rows = check_selection_rationality(prep)
    assert len(rows) == 1   # 1 个触发学生（专业选修 应修 8 差 6）
    path = exp.build_report_xlsx(prep, rows, "2023级")
    assert os.path.exists(path)
    assert "2023级" in os.path.basename(path) and path.endswith(".xlsx")
    import openpyxl
    wb = openpyxl.load_workbook(path)
    try:
        assert any("汇总" in s for s in wb.sheetnames)
        assert wb["汇总"]["A1"].value == "学号"
        assert wb["汇总"]["A2"].value == "S1"
    finally:
        wb.close()


# ---- Task5：导出明细与豁免口径一致（类型B 叠加 + 虚拟行呈现核对）----

def _waiver_prep(waiver_credit=None, virtual_row=False):
    """导出口径测试共用：S1 无真实成绩（豁免作用对象）+ S2 真实修满 EL1/EL2
    （专业覆盖/全员无成绩豁免均需真实成绩行；R1 虚拟行不计覆盖）。"""
    from academicwarning.models import Student, Grade, PlanCourse
    from academicwarning.parsers import clean_course_name
    from academicwarning.rules import Prep
    students = [Student(student_id="S1", name="甲", grade="2023级", major="工商管理"),
                Student(student_id="S2", name="乙", grade="2023级", major="工商管理")]
    courses = [PlanCourse(plan_id=1, course_code="EL1", course_name="选修课A",
                          credit=2.0, course_type="专业选修课程", required_flag="选修",
                          semester="3-1"),
               PlanCourse(plan_id=1, course_code="EL2", course_name="选修课B",
                          credit=2.0, course_type="专业选修课程", required_flag="选修",
                          semester="3-2")]
    grades = [_grade("S2", "选修课A", 2.0), _grade("S2", "选修课B", 2.0)]
    if virtual_row:   # 类型A 注入形态（service _apply_waivers 同款）
        grades.append(Grade(student_id="S1", course_name="选修课A", credit=2.0,
                            grade_raw="特殊豁免", pass_flag=1,
                            course_name_clean=clean_course_name("选修课A")))
    return Prep(students, [], grades, {"工商管理": courses}, {}, [], "4-1",
                entry_year=2023, waiver_credit=waiver_credit)


def test_export_sheets_waiver_credit_gap_consistent(tmp_path, monkeypatch):
    """Sheet2/3 重算 gained 叠加类型B waiver_credit（N3 同款 3 行）——
    已修 2.0、差额 4-2=2.0 与检查 rows/message 数字同口径（Task2 review：
    叠加前 Sheet2/3 差 4.0 与 Sheet1/检查结果不一致）；对照无认可分差 4.0。"""
    from academicwarning.selection_check import check_selection_rationality
    import academicwarning.export as exp
    import openpyxl
    monkeypatch.setattr(exp, "OUT_DIR", str(tmp_path))
    prep = _waiver_prep(waiver_credit={"S1": {"专业选修课程": 2.0}})
    rows = check_selection_rationality(prep)
    assert len(rows) == 1 and rows[0].gap == 2.0
    assert "已修2.0" in rows[0].message        # 检查口径已含认可分
    assert check_selection_rationality(_waiver_prep())[0].gap == 4.0   # 基线差 4
    wb = openpyxl.load_workbook(exp.build_report_xlsx(prep, rows, "2023级"))
    try:
        s2 = next(r for r in wb["学生类别明细"].iter_rows(min_row=2, values_only=True)
                  if r[0] == "S1" and r[3] == "专业选修课程")
        assert s2[4:8] == (4.0, 2.0, 0.0, 2.0)   # 应修/已修/已选/差额
        s3 = next(r for r in wb["全类别总览"].iter_rows(min_row=2, values_only=True)
                  if r[0] == "S1" and r[3] == "专业选修课程")
        assert s3[4:8] == (4.0, 2.0, 0.0, 2.0)
    finally:
        wb.close()


def test_export_sheet2_virtual_row_waiver_mark_shown(tmp_path, monkeypatch):
    """虚拟行（类型A，grade_raw='特殊豁免'）在已修课程明细自然呈现
    '选修课A(2.0学分/特殊豁免)'（L2 标记口径，无需过滤/改写）且计入已修
    （与检查 gained 同源）——差额 = 应修4 − 虚拟行2 = 2.0。"""
    from academicwarning.selection_check import check_selection_rationality
    import academicwarning.export as exp
    import openpyxl
    monkeypatch.setattr(exp, "OUT_DIR", str(tmp_path))
    prep = _waiver_prep(virtual_row=True)
    rows = check_selection_rationality(prep)
    assert len(rows) == 1 and rows[0].gap == 2.0   # 虚拟行计入已修后才触发
    wb = openpyxl.load_workbook(exp.build_report_xlsx(prep, rows, "2023级"))
    try:
        s2 = next(r for r in wb["学生类别明细"].iter_rows(min_row=2, values_only=True)
                  if r[0] == "S1" and r[3] == "专业选修课程")
        assert s2[4:8] == (4.0, 2.0, 0.0, 2.0)
        assert s2[8] == "选修课A(2.0学分/特殊豁免)"
    finally:
        wb.close()


# ---- C4：导出口径同步——标红 差额>0 + 表头"已修"注释注明及格制 ----

def test_export_red_fill_gap_gt_zero_and_header_comment(tmp_path, monkeypatch):
    """C4：Sheet2/3 标红条件由 差额≥2 改 >0——差 1.0 行标红（旧口径不红）、
    同学生差额 0 的类别行不红；表头"已修"单元格带口径注释（author 非空）。
    装配参照 test_selection_check.test_gap_threshold：S1 只修 EL1(2.0)、
    应修 EL1+EL2(1.0)=3.0 → 差 1.0 触发；S2 全修防全员无成绩豁免。"""
    from academicwarning.models import Student, PlanCourse
    from academicwarning.rules import Prep
    from academicwarning.selection_check import check_selection_rationality
    import academicwarning.export as exp
    import openpyxl
    monkeypatch.setattr(exp, "OUT_DIR", str(tmp_path))   # 落盘临时目录
    courses = [PlanCourse(plan_id=1, course_code="EL1", course_name="选修课A",
                          credit=2.0, course_type="专业选修课程", required_flag="选修",
                          semester="3-1"),
               PlanCourse(plan_id=1, course_code="EL2", course_name="选修课B",
                          credit=1.0, course_type="专业选修课程", required_flag="选修",
                          semester="3-2")]
    students = [Student(student_id="S1", name="甲", grade="2023级", major="工商管理"),
                Student(student_id="S2", name="乙", grade="2023级", major="工商管理")]
    grades = [_grade("S1", "选修课A", 2.0),   # S1 只过 EL1 → 差 1.0
              _grade("S2", "选修课A", 2.0),   # S2 全修 → 应累计不豁免
              _grade("S2", "选修课B", 1.0)]
    prep = Prep(students, [], grades, {"工商管理": courses}, {}, [], "4-1",
                entry_year=2023)
    rows = check_selection_rationality(prep)
    assert len(rows) == 1 and rows[0].gap == 1.0
    wb = openpyxl.load_workbook(exp.build_report_xlsx(prep, rows, "2023级"))
    try:
        for name in ("学生类别明细", "全类别总览"):
            ws = wb[name]
            r = next(r for r in ws.iter_rows(min_row=2)
                     if r[0].value == "S1" and r[3].value == "专业选修课程")
            assert r[7].value == 1.0                    # 差额 1.0 → 触发标红
            assert r[7].fill.patternType == "solid"
            assert r[7].fill.start_color.rgb.endswith("FDE9E9")
            c = ws["F1"]                                # 表头"已修"
            assert c.value == "已修"
            assert c.comment is not None and c.comment.author
            assert "及格制" in c.comment.text and "明细" in c.comment.text
        # Sheet3 每类别一行：同学生差额 0.0 的类别行不标红
        s3 = wb["全类别总览"]
        zero = next(r for r in s3.iter_rows(min_row=2)
                    if r[0].value == "S1" and r[3].value == "模块课程")
        assert zero[7].value == 0.0
        assert zero[7].fill.patternType is None
    finally:
        wb.close()


# ---- final review F1/F2：判重按年级分区 + export 年级参数校验 ----

def test_dedup_partitioned_by_grade(tmp_path):
    """F1：判重按年级分区——2026 级复用 2023 级同内容方案不判重（可注册到
    第二年级）；同年级内同 hash 仍判重拦截。"""
    from datetime import datetime
    from academicwarning.service import dedup_check
    db = _db(tmp_path)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    def _plan(g):
        db.insert_source_file(SourceFile(
            file_type="plan", file_name=f"方案-{g}.docx", file_hash="SAME",
            file_path="/tmp/x", upload_time=now, uploader="t", grade=g,
            in_file_meta={"major": "工商管理"}))

    _plan("2023级")
    # 跨年级：同 hash 不同 grade 不判重（未命中 → 继续入库）
    assert db.file_hash_exists("SAME", "plan", "2023级") is not None
    assert db.file_hash_exists("SAME", "plan", "2026级") is None
    assert dedup_check(db, "SAME", "plan", "2026级") is None
    assert "已上传过" in dedup_check(db, "SAME", "plan", "2023级")
    # 同年级：第二份同 hash 文件仍判重拦截
    _plan("2023级")
    assert "已上传过" in dedup_check(db, "SAME", "plan", "2023级")


def test_export_grade_format_validated(client, api_env):
    """F2：export 端点 grade 参数格式校验——非 20xx级 → 404（年级参数无效），
    合法格式走原有路径（无报告 → 暂无导出的检查报告）。"""
    r = client.get("/api/warning/selection-check/export?grade=23级",
                   headers={"X-API-Key": TEST_KEY})
    assert r.status_code == 404
    assert r.json()["detail"] == "年级参数无效"
    # 合法格式（2099级 符合 20xx级 但必无报告）→ 走原有 404 路径，不被格式校验拦截
    r2 = client.get("/api/warning/selection-check/export?grade=2099级",
                    headers={"X-API-Key": TEST_KEY})
    assert r2.status_code == 404
    assert "暂无导出" in r2.json()["detail"]


# ---- v6.3（2026-09-02）：年级当前学期配置 + 成绩覆盖自动判定/人工确认 + 豁免分级 ----

def test_term_label_to_sem():
    """term_label → 学期编码（'第X学年（YYYY-YYYY）第Y学期'，三=小学期）。"""
    from academicwarning.parsers import term_label_to_sem
    assert term_label_to_sem("第三学年（2025-2026）第二学期") == "3-2"
    assert term_label_to_sem("第二学年（2024-2025）第三学期") == "2-3"
    assert term_label_to_sem("第一学年（2023-2024）第一学期") == "1-1"
    assert term_label_to_sem("乱数据") == ""
    # 错位标签（污染数据）：'第二学年（2023-2024）第二学期' → 2-2（年份 2023-2024
    # 与学年 2 矛盾——取学年序号为准；覆盖取 max 时此类错位值 ≤ 正常值无影响）
    assert term_label_to_sem("第二学年（2023-2024）第二学期") == "2-2"
    # 休学复学生（2026-09-02 v6.4 放宽）：学年一~六、半角括号与空白、尾部
    # [休学] 等标注均可解析（如 2021 级休学复学生最长到第五/六学年）
    assert term_label_to_sem("第五学年（2025-2026）第二学期") == "5-2"
    assert term_label_to_sem("第六学年（2026-2027）第一学期") == "6-1"
    assert term_label_to_sem("第一学年 (2021-2022) 第二学期 [休学]") == "1-2"
    assert term_label_to_sem("第二学年 (2022-2023) 第三学期 [复学]") == "2-3"


def _pc(code, name, semester, credit=2.0):
    from academicwarning.models import PlanCourse
    return PlanCourse(plan_id=1, course_code=code, course_name=name,
                      credit=credit, course_type="集中实践",
                      required_flag="必修", semester=semester)


def _prep41(courses, sels, covered_sem="", override=""):
    """分级豁免测试用 Prep：工商双学生（成绩单各 1 条高数，专业有覆盖非空）。"""
    from academicwarning.models import Student
    from academicwarning.rules import Prep
    students = [Student(student_id="S1", name="甲", grade="2023级", major="工商管理"),
                Student(student_id="S2", name="乙", grade="2023级", major="工商管理")]
    return Prep(students, sels, [_grade("S1", "高等数学", 4.0),
                                 _grade("S2", "高等数学", 4.0)],
                {"工商管理": courses}, {}, [], "4-1", entry_year=2023,
                covered_sem=covered_sem, current_sem_override=override)


def test_exempt_graded_by_coverage():
    """分级豁免（规则 B + 覆盖学期）：≤covered(3-2) 无成绩 = 真缺 → 不豁免；
    3-3 未出窗口 → 豁免；4-1 有人选 → 不豁免；4-1 无人选 → 豁免；复合学期 → 豁免。"""
    courses = [_pc("C1", "已出学期必修A", "1-1"),
               _pc("C2", "小学期实习B", "3-3"),
               _pc("C3", "本学期开课C", "4-1"),
               _pc("C4", "讲座研讨D", "4-1"),
               _pc("C5", "形势与政策E", "1-1至4-1")]
    prep = _prep41(courses, [_sel("S1", "C3", "本学期开课C")],
                   covered_sem="3-2", override="4-1")
    # C2/C4/C5 豁免；C1 成绩应已出无记录=真缺、C3 当前学期有人选=规则 B → 不豁免
    assert prep.missing_courses("S1") == {"C2", "C4", "C5"}
    # 缺修必修逐人视角：S1 只缺 C1（C3 已选差 0）；S2 未选 C3 → 缺 C1+C3
    from academicwarning.selection_check import _missing_required_courses
    assert [n for _, n, _, _, _, _ in _missing_required_courses(prep, "S1")] == ["已出学期必修A"]
    assert [n for _, n, _, _, _, _ in _missing_required_courses(prep, "S2")] == [
        "已出学期必修A", "本学期开课C"]


def test_exempt_covered_empty_default():
    """covered_sem 默认 ""（既有调用不传）退化正确：判定 2（≤covered）恒 False →
    全员无成绩豁免恢复 v6.2 行为；规则 B（当前学期有人选）仍生效。"""
    courses = [_pc("C1", "已出学期必修A", "1-1"),
               _pc("C2", "小学期实习B", "3-3"),
               _pc("C3", "本学期开课C", "4-1"),
               _pc("C4", "讲座研讨D", "4-1")]
    prep = _prep41(courses, [])
    assert prep.missing_courses("S1") == {"C1", "C2", "C3", "C4"}
    prep2 = _prep41(courses, [_sel("S1", "C3", "本学期开课C")])
    assert prep2.missing_courses("S1") == {"C1", "C2", "C4"}


def test_config_semester_override_file_inference():
    """current_sem_override 优先：文件推断 4-1 vs 配置 3-1 → 应修只到 3-1。"""
    courses = [_pc("C0", "大三课程", "3-1"), _pc("C1", "大四课程", "4-1")]
    prep_cfg = _prep41(courses, [], override="3-1")
    assert prep_cfg.semester_code == "3-1"
    assert {c.course_code for c in prep_cfg._due_courses_by_major["工商管理"]} == {"C0"}
    assert "C1" not in prep_cfg.missing_courses("S1")   # 4-1 未开课不参与
    prep_file = _prep41(courses, [])
    assert prep_file.semester_code == "4-1"             # 无配置 → 文件推断
    assert {c.course_code for c in prep_file._due_courses_by_major["工商管理"]} == {"C0", "C1"}
    assert "C1" in prep_file.missing_courses("S1")


def test_grade_covered_semester_normalized(tmp_path):
    """成绩覆盖自动解析（I1 修复）：auto 与 confirm 同口径「整批」——某学期须
    同一专业 ≥2 名学生（归一后）含记录才计入该专业候选；跨专业最大学期仍取
    最小（全年级同步下发口径）。单专业零星 1 人 3-3 / 单专业 2 人但其余专业
    未同步 → auto 不升；四专业同批各 ≥2 人含 3-3 → auto 升 3-3。"""
    from academicwarning.models import Grade
    from academicwarning.service import grade_covered_semester
    db = _db(tmp_path)
    t32 = "第三学年（2025-2026）第二学期"
    t33 = "第三学年（2025-2026）第三学期"
    majors = ["工商管理", "大数据管理与应用", "工业工程", "会计学（ACCA）"]
    seq = [0]

    def _gfile(major, rows):
        """rows: [(term_label, 学生id), ...]；后插入 = 该专业最新文件。"""
        seq[0] += 1
        gid = db.insert_source_file(SourceFile(
            file_type="grade", file_name=f"g-{seq[0]}.docx", file_hash=f"hg-{seq[0]}",
            file_path="", upload_time="2026-09-01 00:00:00", uploader="t",
            parsed_status="done", grade="2023级", in_file_meta={"major": major}))
        db.insert_grades([Grade(student_id=sid, student_name=f"学{sid}",
                                term_label=lab, course_name="高等数学", credit=4.0,
                                grade_raw="80", source_file_id=gid)
                          for lab, sid in rows])
        return gid

    # 全年级 3-2 整批（各专业 ≥2 名学生）
    for m in majors:
        _gfile(m, [(t32, f"{m}A"), (t32, f"{m}B")])
    # 工商/大数据下一次文件各带 1 条零星 3-3（补修/复学页，2026-09 现状）→ auto 不升
    _gfile("工商管理", [(t32, "GS-A"), (t32, "GS-B"), (t33, "GS-X")])
    _gfile("大数据管理与应用", [(t32, "DS-A"), (t32, "DS-B"), (t33, "DS-X")])
    assert grade_covered_semester(db, "2023级") == ("3-2", "3-2")
    # F1 判别形态：工业/ACCA 最新文件也各带 1 条零星 3-3 → 四专业全部只有零星
    # 3-3（无整批）→ auto 仍 3-2（修复前 ≥1 行实现会错升 3-3，此断言为判别断言）
    _gfile("工业工程", [(t32, "GY-A"), (t32, "GY-B"), (t33, "GY-X")])
    _gfile("会计学（ACCA）", [(t32, "AC-A"), (t32, "AC-B"), (t33, "AC-X")])
    assert grade_covered_semester(db, "2023级") == ("3-2", "3-2")
    # 错位标签（第四学年（2025-2026）= 2022 届口径）与半角括号污染 → 不影响
    gs_latest = db.latest_source_files_by_major("grade", "2023级")["工商管理"]
    db.insert_grades([
        Grade(student_id="GS-Y", student_name="跨届生",
              term_label="第四学年（2025-2026）第一学期",
              course_name="高年级课", credit=2.0, grade_raw="80",
              source_file_id=gs_latest.id),
        Grade(student_id="GS-Y", student_name="跨届生",
              term_label="第一学年 (2021-2022) 第二学期 [休学]",
              course_name="旧课", credit=2.0, grade_raw="80",
              source_file_id=gs_latest.id)])
    assert grade_covered_semester(db, "2023级") == ("3-2", "3-2")
    # 单专业达 2 名学生 3-3（整批口径达标），其余专业未同步 → min 口径仍不升
    _gfile("工商管理", [(t32, "GS-C"), (t32, "GS-D"), (t33, "GS-E"), (t33, "GS-F")])
    assert grade_covered_semester(db, "2023级") == ("3-2", "3-2")
    # 四专业成绩单同批含 3-3（各 ≥2 名学生，9-10 月下发）→ auto 自动升 3-3
    for m, p in zip(majors, ("P", "Q", "R", "S"), strict=False):
        _gfile(m, [(t33, f"{p}1"), (t33, f"{p}2")])
    assert grade_covered_semester(db, "2023级") == ("3-3", "3-3")


def test_update_and_confirm_grade_semester(client, api_env):
    """PUT /grade 校验 11 项学期/不存在年级；confirm 校验成绩单整批记录（≥2 名学生
    + 年份归一），confirm_sem 落库并抬升 covered。"""
    from academicwarning.models import Grade
    from academicwarning.service import grade_covered_semester
    majors = ["工商管理", "大数据管理与应用", "工业工程", "会计学（ACCA）"]
    db = WarningDB(api_env)
    db.add_grade("2023级")
    db.close()
    # PUT：非法学期（含 4-3 超出学制）拒绝
    for bad in ("9-9", "4-3", "x-y"):
        r = client.put("/api/warning/grade",
                       json={"name": "2023级", "current_semester": bad},
                       headers={"X-API-Key": TEST_KEY})
        assert r.json()["ok"] is False and "学期无效" in r.json()["message"]
    # 不存在的年级拒绝
    r = client.put("/api/warning/grade",
                   json={"name": "2024级", "current_semester": "3-1"},
                   headers={"X-API-Key": TEST_KEY})
    assert r.json()["ok"] is False and "不存在" in r.json()["message"]
    # 合法设置 + db 读回
    r = client.put("/api/warning/grade",
                   json={"name": "2023级", "current_semester": "4-1"},
                   headers={"X-API-Key": TEST_KEY})
    assert r.json()["ok"] is True
    db = WarningDB(api_env)
    assert db.get_grade_semesters("2023级") == ("4-1", "")
    # 空 = 恢复自动推断
    r = client.put("/api/warning/grade",
                   json={"name": "2023级", "current_semester": ""},
                   headers={"X-API-Key": TEST_KEY})
    assert r.json()["ok"] is True
    assert db.get_grade_semesters("2023级") == ("", "")
    db.set_grade_semester("2023级", "4-1")
    # confirm：无成绩单记录 → 拒绝
    r = client.post("/api/warning/grade/confirm",
                    json={"name": "2023级", "sem": "3-2"},
                    headers={"X-API-Key": TEST_KEY})
    assert r.json()["ok"] is False and "未检测到" in r.json()["message"]

    # 四专业成绩单均含 3-2（整批）→ auto=3-2
    t32, t33 = "第三学年（2025-2026）第二学期", "第三学年（2025-2026）第三学期"
    for m in majors:
        gid = db.insert_source_file(SourceFile(
            file_type="grade", file_name=f"g-{m}.docx", file_hash=f"h-{m}",
            file_path="", upload_time="2026-09-01 00:00:00", uploader="t",
            parsed_status="done", grade="2023级", in_file_meta={"major": m}))
        db.insert_grades([
            Grade(student_id="S1", student_name="甲", term_label=t32,
                  course_name="高等数学", credit=4.0, grade_raw="80",
                  source_file_id=gid),
            Grade(student_id="S2", student_name="乙", term_label=t32,
                  course_name="高等数学", credit=4.0, grade_raw="90",
                  source_file_id=gid)])
    assert grade_covered_semester(db, "2023级") == ("3-2", "3-2")
    # confirm 3-2 ≤ auto → 幂等成功（自动检测已含）
    r = client.post("/api/warning/grade/confirm",
                    json={"name": "2023级", "sem": "3-2"},
                    headers={"X-API-Key": TEST_KEY})
    assert r.json()["ok"] is True and "自动检测已含" in r.json()["message"]
    assert db.get_grade_semesters("2023级") == ("4-1", "3-2")
    # confirm 3-3：仅工商 1 名补修生有记录（2026-09 现状）→ 拒绝（非整批）
    gs = db.latest_source_files_by_major("grade", "2023级")["工商管理"]
    db.insert_grades([Grade(student_id="S9", student_name="补修生", term_label=t33,
                            course_name="大学计算机 - 算法编程", credit=3.0,
                            grade_raw="60", source_file_id=gs.id)])
    r = client.post("/api/warning/grade/confirm",
                    json={"name": "2023级", "sem": "3-3"},
                    headers={"X-API-Key": TEST_KEY})
    assert r.json()["ok"] is False and "未检测到" in r.json()["message"]
    # 同专业第 2 名学生也有 3-3 记录 → 确认成功（confirm_sem 落库、covered 抬升）
    db.insert_grades([Grade(student_id="S2", student_name="乙", term_label=t33,
                            course_name="专业实习II", credit=4.0, grade_raw="85",
                            source_file_id=gs.id)])
    r = client.post("/api/warning/grade/confirm",
                    json={"name": "2023级", "sem": "3-3"},
                    headers={"X-API-Key": TEST_KEY})
    assert r.json()["ok"] is True
    assert db.get_grade_semesters("2023级") == ("4-1", "3-3")
    assert grade_covered_semester(db, "2023级") == ("3-2", "3-3")
    # 学期格式无效 / 晚于当前学期 → 拒绝
    r = client.post("/api/warning/grade/confirm",
                    json={"name": "2023级", "sem": "4-3"},
                    headers={"X-API-Key": TEST_KEY})
    assert r.json()["ok"] is False and "学期格式无效" in r.json()["message"]
    db.set_grade_semester("2023级", "3-1")
    r = client.post("/api/warning/grade/confirm",
                    json={"name": "2023级", "sem": "3-2"},
                    headers={"X-API-Key": TEST_KEY})
    assert r.json()["ok"] is False and "晚于当前学期" in r.json()["message"]
    db.close()


def test_selection_check_response_sem_fields(client, api_env):
    """GET /selection-check 响应带 current_semester/covered_sem/covered_auto/confirm_sem
    （前端状态行数据源）；无年级 = 旧客户端 → 空值。"""
    db = WarningDB(api_env)
    db.add_grade("2023级")
    db.set_grade_semester("2023级", "4-1")
    db.close()
    r = client.get("/api/warning/selection-check?grade=2023级",
                   headers={"X-API-Key": TEST_KEY})
    assert r.status_code == 200
    d = r.json()
    assert d["current_semester"] == "4-1" and d["confirm_sem"] == ""
    assert d["covered_sem"] == "" and d["covered_auto"] == ""
    r2 = client.get("/api/warning/selection-check", headers={"X-API-Key": TEST_KEY})
    assert r2.status_code == 200
    assert r2.json()["current_semester"] == ""
    assert r2.json()["covered_sem"] == "" and r2.json()["covered_auto"] == ""


# ---- 2026-09-07：通识表全局单份上传（gen_ed source_file；上传优先、删除回落内置）----
def test_gen_ed_upload_global_and_loader_prefers_uploaded(api_env, tmp_path, monkeypatch):
    """上传通识表（不带年级=全局单份）→ source_file done、loader(db) 取上传版学分；
    状态区在任意年级过滤下仍可见该行。"""
    import academicwarning.service as svc
    monkeypatch.setattr(svc, "_GEN_ED_TABLE", None)
    monkeypatch.setattr(svc, "_GEN_ED_TABLE_KEY", None)
    f = tmp_path / "通识课程信息表.xlsx"
    _write_gen_ed_table(f, [("GNED999001", "全球科技史", 4.0)])
    db = WarningDB(api_env)
    client = TestClient(api.app)
    with open(f, "rb") as fh:
        r = client.post(
            "/api/warning/upload",
            files={"file": ("通识课程信息表.xlsx", fh, "application/octet-stream")},
            data={"type_hint": "通识课程信息表", "grade_hint": ""},
            headers={"X-API-Key": TEST_KEY})
    assert r.status_code == 200
    assert r.json()["parsed_status"] in ("parsing", "done")
    deadline = time.time() + 15
    sf = None
    while time.time() < deadline:
        sf = db.latest_source_file("gen_ed", "")
        if sf and sf.parsed_status == "done":
            break
        time.sleep(0.2)
    assert sf is not None and sf.parsed_status == "done"
    assert sf.grade == ""   # 全局单份不落年级
    by_code, _ = svc._load_gen_ed_table(db)
    assert by_code.get("GNED999001") == 4.0   # 上传版优先于内置默认表
    # 状态区在年级过滤下仍含 gen_ed 行
    st = client.get("/api/warning/status?grade=2023级",
                    headers={"X-API-Key": TEST_KEY}).json()
    ge = next(g for g in st["files"] if g["file_type"] == "gen_ed")
    assert ge["items"][0]["file_name"] == "通识课程信息表.xlsx"
    db.close()


def test_gen_ed_delete_ignores_grade_param(api_env, tmp_path, monkeypatch):
    """前端删除带 grade 参数（全局文件无年级）→ 仍删全局记录并回落内置表。"""
    import academicwarning.service as svc
    monkeypatch.setattr(svc, "_GEN_ED_TABLE", None)
    monkeypatch.setattr(svc, "_GEN_ED_TABLE_KEY", None)
    f = tmp_path / "通识课程信息表.xlsx"
    _write_gen_ed_table(f, [("GNED999002", "全球艺术史", 1.0)])
    db = WarningDB(api_env)
    client = TestClient(api.app)
    with open(f, "rb") as fh:
        client.post(
            "/api/warning/upload",
            files={"file": ("通识课程信息表.xlsx", fh, "application/octet-stream")},
            data={"type_hint": "通识课程信息表"},
            headers={"X-API-Key": TEST_KEY})
    deadline = time.time() + 15
    while time.time() < deadline:
        sf = db.latest_source_file("gen_ed", "")
        if sf and sf.parsed_status == "done":
            break
        time.sleep(0.2)
    assert sf is not None and sf.parsed_status == "done"
    r = client.post("/api/warning/delete",
                    json={"file_type": "gen_ed", "grade": "2023级"},
                    headers={"X-API-Key": TEST_KEY})
    assert r.status_code == 200
    assert "暂无" not in r.json()["message"]
    assert db.latest_source_file("gen_ed", "") is None   # 全局记录已删
    db.close()


@pytest.mark.samples
def test_gen_ed_delete_falls_back_to_default(api_env, tmp_path, monkeypatch):
    """删除上传的通识表 → loader(db) 回落内置默认表（docs 静态）。

    2026-09-22：内置表 `academicwarning/docs/通识课程信息表.xlsx` 是**代码资产**，
    但该目录被 agent4som/.gitignore 整体排除（内含真实学生数据），且经用户决定
    **不入库**（只保留运行时上传版）。故本测试在内置表缺失时 skip——回落到
    ({}, {}) 是预期行为，不算失败。"""
    import academicwarning.service as svc
    if not svc._GEN_ED_TABLE_FILE.exists():
        pytest.skip("内置通识表 docs/通识课程信息表.xlsx 缺失（用户决定不入库）——"
                    "兜底路径回落空表，见 HANDOFF 技术债 #10")
    monkeypatch.setattr(svc, "_GEN_ED_TABLE", None)
    monkeypatch.setattr(svc, "_GEN_ED_TABLE_KEY", None)
    db = WarningDB(api_env)
    from academicwarning.db import SourceFile as SF
    now = "2026-09-07 10:00:00"
    fid = db.insert_source_file(SF(
        file_type="gen_ed", file_name="自定义表.xlsx", file_hash="h9",
        file_path=str(tmp_path / "自定义表.xlsx"), upload_time=now,
        uploader="t", parsed_status="done", grade="",
        in_file_meta={"rows": 1}))
    (tmp_path / "自定义表.xlsx").write_bytes(b"x")   # 占位（loader 失败会回落默认）
    by_code, _ = svc._load_gen_ed_table(db)
    assert by_code == {} or any(k.startswith("CORE") or k.startswith("GNED")
                                for k in by_code)   # 文件损坏 → 回落内置默认表
    db.delete_major_files("gen_ed")
    svc._GEN_ED_TABLE = None
    svc._GEN_ED_TABLE_KEY = None
    by_code2, _ = svc._load_gen_ed_table(db)
    assert len(by_code2) > 400   # 内置 docs 表回落生效
    db.close()
