"""成绩单 parse_grades 合成 docx 测试（此前零本地执行覆盖——真实样本
academicwarning/docs/*.docx 不随仓库分发，test_parsers_grade.py 相关用例全 skip）。

内存构造 docx（图片 + 表格按 §2 关联规则2 一一对应），vl_func 用注入 fake
（按图片纵横比分流：宽高比 <5 的图返回学号文本，≥5 恒空串模拟 OCR 失败），
覆盖：OCR 失败计入 failed_idx、孤儿课/节序回退报备输出、图片数≠表格数告警。
"""
import io

from docx import Document
from PIL import Image

import academicwarning.parsers as ap
from academicwarning.parsers import parse_grades

_T1 = "第一学年（2023-2024）第一学期"
_T2 = "第二学年（2024-2025）第二学期"
_TR = "第一学年（2023-2024）第二学期"   # 节序回退（2-2 后出 1-2）


def _png(w: int, h: int) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (w, h), "white").save(buf, format="PNG")
    return buf.getvalue()


def _vl_by_aspect(path: str) -> str:
    """按纵横比分流的 fake VL：窄图（纵横比 <5）→ 学号文本；宽条图 → 恒空串。

    OCR 三策略分别 4x/1x/3x 缩放，均保持纵横比，分流对策略不敏感。"""
    w, h = Image.open(path).size
    if w / h < 5:
        return "学号：S2233001\n姓名：测试甲"
    return ""


def _grade_doc():
    doc = Document()
    # 图1（12x6，纵横比 2）→ OCR 成功；表1
    doc.add_picture(io.BytesIO(_png(12, 6)))
    t1 = doc.add_table(rows=9, cols=3)
    rows = [
        ["成绩单", "", ""],               # R0 标题行（rows[0] 不读）
        ["课程", "学分", "成绩"],          # R1 表头
        ["孤儿课A", "1.0", "60"],          # R2 首节之前的课程 → 孤儿报备
        [_T1, "", ""],                     # R3 1-1 节
        ["高数", "4.0", "80"],             # R4
        [_T2, "", ""],                     # R5 2-2 节
        ["大物", "3.0", "70"],             # R6
        [_TR, "", ""],                     # R7 节序回退 → 报备
        ["课后课", "1.0", "90"],           # R8 按回退节继续解析
    ]
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t1.cell(ri, ci).text = val
    # 图2（60x6，纵横比 10）→ OCR 恒失败；表2（内容不解析）
    doc.add_picture(io.BytesIO(_png(60, 6)))
    t2 = doc.add_table(rows=3, cols=3)
    for ri, row in enumerate([["成绩单", "", ""], ["课程", "学分", "成绩"],
                              [_T1, "", ""]]):
        for ci, val in enumerate(row):
            t2.cell(ri, ci).text = val
    # 图3（无对应表格）→ 图片数 ≠ 表格数告警
    doc.add_picture(io.BytesIO(_png(8, 8)))
    return doc


def test_parse_grades_synth(tmp_path, caplog, monkeypatch):
    monkeypatch.setattr(ap, "_OCR_RETRY_SLEEP", 0)   # OCR 失败补跑间隔置 0
    path = tmp_path / "成绩单.docx"
    _grade_doc().save(path)
    students, failed_idx = parse_grades(str(path), vl_func=_vl_by_aspect)

    # 表1 OCR 成功解析，表2 失败 → failed_idx
    assert failed_idx == [1]
    assert len(students) == 1
    sid, name, grades = students[0]
    assert (sid, name) == ("S2233001", "测试甲")
    assert [(g.course_name, g.term_label) for g in grades] == [
        ("孤儿课A", ""), ("高数", _T1), ("大物", _T2), ("课后课", _TR)]
    # 报备输出（logging）：告警 / 孤儿课 / 节序回退
    text = caplog.text
    assert "图片 3 张 ≠ 成绩单 2 份" in text
    assert "孤儿课：1 门课程" in text and "表#0（1 门）：孤儿课A" in text
    assert "节序异常：1 份成绩单" in text and f"[{_T2}] 之后回退到 [{_TR}]" in text
