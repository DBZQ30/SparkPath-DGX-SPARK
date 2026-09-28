"""培养方案 parse_plan 合成 docx 测试。

真实样本（academicwarning/docs/*.docx）不随仓库分发（tests/academicwarning/
test_parsers_plan.py 全 skip），据此用 python-docx 内存构造覆盖各分支：
表头定位（rows[0]/rows[1]）、课程组子名修复、无效行过滤、必修选修取词、
学期格式校验、推荐课表跨列合并表头、毕业要求学分表、专业名标题/文件名兜底。
"""
import pytest
from docx import Document

from academicwarning.parsers import parse_plan

# 课程总表列：课程类型|课程编码|中文课程名称|中文课程名称|学分|必修选修|开课学期|开课单位
_HDR = ["课程类型", "课程编码", "中文课程名称", "中文课程名称",
        "学分", "必修选修", "开课学期", "开课单位"]
# 合并列模拟：python-docx 读合并单元格会在各列重复同一文本，直接重复填充即可
_SEM_HDR = ["第一学期：1-1", "第一学期：1-1", "第一学期：1-1",
            "第二学期：1-2", "第二学期：1-2", "第二学期：1-2"]


def _fill(table, rows):
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.cell(ri, ci).text = val


def _full_doc():
    """全要素方案：学分结构表 + 课程总表（含各类无效行）+ 推荐课表。"""
    doc = Document()
    doc.add_paragraph("工业工程专业培养方案")
    # Table 0：学分结构要求表
    t0 = doc.add_table(rows=2, cols=3)
    _fill(t0, [["课程类别", "毕业要求学分", "备注"],
               ["专业选修课程", "8", "含跨专业选修"]])
    # Table 1：课程总表
    t1 = doc.add_table(rows=10, cols=8)
    _fill(t1, [
        _HDR,
        ["专业核心课程", "REQ1", "必修高数A", "必修高数A", "4", "必修", "2-1", "数学学院"],
        # 课程组：名称占 2 列，首列组名、第二列子名（成绩单记录的是子名）
        ["专业选修课程", "EL1", "思想政治理论", "思想道德与法治", "3", "必修\n15学分", "1-1", "马院"],
        ["小计", "", "", "", "40", "", "", ""],                     # 小计行剔除
        ["神秘类型", "X1", "未知课", "未知课", "2", "必修", "1-1", "x"],   # 未知类型剔除
        ["", "X2", "空类型课", "空类型课", "2", "必修", "1-1", "x"],      # 空类型剔除
        ["专业核心课程", "REQ2", "劳动教育", "劳动教育", "0", "必修", "1-1", "x"],  # 0 学分剔除
        ["专业核心课程", "REQ3", "四选一体育", "四选一体育", "2", "必修", "", "x"],  # 无学期剔除
        ["集中实践", "PRAC1", "毕业论文", "毕业论文", "6", "选修", "4-2", "商学院"],
        # 学期格式异常：计入 bad_semesters，课程行仍保留
        ["专业选修课程", "EL9", "创新实践", "创新实践", "2", "—", "第4学年", "x"],
    ])
    # Table 2：推荐课表（数据从 rows[2:] 起；含合计行）
    t2 = doc.add_table(rows=4, cols=6)
    _fill(t2, [
        _SEM_HDR,
        ["（表头延续行，数据自 rows[2]）"] + [""] * 5,
        ["REQ1", "必修高数A", "4", "EL1", "思想道德与法治", "3"],
        ["合计", "总学分", "22", "", "", ""],
    ])
    return doc


def test_parse_plan_synth_full(tmp_path):
    """全要素合成方案：有效行解析、无效行过滤、组名子名修复、必修取词
    （含默认必修）、学期格式异常清单、推荐课表与毕业要求学分。"""
    path = tmp_path / "2023版工业工程专业培养方案.docx"
    _full_doc().save(path)
    meta, courses, sem_courses, bad = parse_plan(str(path))

    assert meta == {"major": "工业工程", "elective_req": 8.0}
    by_code = {c.course_code: c for c in courses}
    assert set(by_code) == {"REQ1", "EL1", "PRAC1", "EL9"}
    r1 = by_code["REQ1"]
    assert (r1.course_name, r1.credit, r1.required_flag, r1.semester,
            r1.course_type, r1.provider) == \
        ("必修高数A", 4.0, "必修", "2-1", "专业核心课程", "数学学院")
    # 组名修复：取子名；"必修\n15学分" → 必修
    assert by_code["EL1"].course_name == "思想道德与法治"
    assert by_code["EL1"].required_flag == "必修"
    assert by_code["PRAC1"].required_flag == "选修"
    # 必修选修列无效值（"—"）→ 默认必修；异常学期仍入表但计入清单
    assert by_code["EL9"].required_flag == "必修"
    assert bad == ["创新实践(第4学年)"]
    # 推荐课表：每学期首列表头解析（1-1 / 1-2），合计行剔除
    assert [(s.semester, s.course_code, s.course_name, s.credit)
            for s in sem_courses] == \
        [("1-1", "REQ1", "必修高数A", 4.0), ("1-2", "EL1", "思想道德与法治", 3.0)]


def test_parse_plan_major_from_filename(tmp_path):
    """正文无标题 → 文件名兜底："2023版X专业培养方案.docx" → X。"""
    doc = Document()
    doc.add_paragraph("（无标题的方案正文）")
    t0 = doc.add_table(rows=5, cols=8)
    _fill(t0, [_HDR] + [["专业核心课程", "REQ1", "必修高数A", "必修高数A",
                         "4", "必修", "2-1", "x"]] * 4)
    path = tmp_path / "2023版工商管理专业培养方案.docx"
    doc.save(path)
    meta, _courses, _sem, _bad = parse_plan(str(path))
    assert meta["major"] == "工商管理"
    assert "elective_req" not in meta   # 无学分结构表 → 不设键


def test_parse_plan_header_at_row1(tmp_path):
    """课程总表表头在 rows[1]（首个数据行前置了说明行）→ 兼容定位。"""
    doc = Document()
    doc.add_paragraph("工业工程专业培养方案")
    doc.add_table(rows=2, cols=3)   # Table 0：无课程编码表头（被跳过）
    t1 = doc.add_table(rows=4, cols=8)
    _fill(t1, [
        ["（课程总表说明行）"] + [""] * 7,
        _HDR,
        ["专业核心课程", "REQ1", "必修高数A", "必修高数A", "4", "必修", "2-1", "x"],
        ["专业核心课程", "REQ2", "必修英语B", "必修英语B", "2", "必修", "1-1", "x"],
    ])
    path = tmp_path / "p.docx"
    doc.save(path)
    _meta, courses, _sem, _bad = parse_plan(str(path))
    assert [c.course_code for c in courses] == ["REQ1", "REQ2"]


def test_parse_plan_no_course_table_asserts(tmp_path):
    """无课程编码表头（无课程总表）→ AssertionError（上游入库前拦截）。"""
    doc = Document()
    doc.add_paragraph("工业工程专业培养方案")
    doc.add_table(rows=3, cols=3)
    path = tmp_path / "p.docx"
    doc.save(path)
    with pytest.raises(AssertionError, match="未找到课程总表"):
        parse_plan(str(path))
