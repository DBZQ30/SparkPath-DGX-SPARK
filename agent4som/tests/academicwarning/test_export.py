"""export.py xlsx 报告导出测试（195 行，此前仅 OUT_DIR 隔离 fixture，无输出内容校验）。

复用 test_selection_check 的 Prep 构造方式跑真实检查，导出后用 openpyxl 读回
验证三个 sheet 的结构与内容；另覆盖 _prune_reports 的保留策略。
"""

from __future__ import annotations

import glob
import os

from openpyxl import load_workbook

from academicwarning.export import _prune_reports
from academicwarning.models import Student, Grade, PlanCourse
from academicwarning.parsers import clean_course_name
from academicwarning.rules import Prep
from academicwarning.selection_check import check_selection_rationality

# conftest._isolate_report_dir（autouse）已把 export.OUT_DIR 指到 tmp_path。


def _grade(sid, name, credit, pass_flag=1):
    return Grade(student_id=sid, course_name=name, credit=credit, grade_raw="80",
                 pass_flag=pass_flag, marker="", course_name_clean=clean_course_name(name))


def _make_prep():
    """双学生：S1 专业选修零已修零已选 → 触发差额；S2 修满（防全员无成绩豁免）。"""
    students = [
        Student(student_id="S1", name="甲", grade="2023级", major="0824工商管理"),
        Student(student_id="S2", name="乙", grade="2023级", major="0824工商管理"),
    ]
    courses = [
        PlanCourse(plan_id=1, course_code="EL1", course_name="选修课A",
                   credit=2.0, course_type="专业选修课程", required_flag="选修", semester="3-1"),
        PlanCourse(plan_id=1, course_code="EL2", course_name="选修课B",
                   credit=2.0, course_type="专业选修课程", required_flag="选修", semester="3-2"),
        PlanCourse(plan_id=1, course_code="CR1", course_name="核心课甲",
                   credit=3.0, course_type="专业核心课程", required_flag="必修", semester="2-1"),
    ]
    grades = [_grade("S2", "选修课A", 2.0), _grade("S2", "选修课B", 2.0),
              _grade("S2", "核心课甲", 3.0)]
    return Prep(students, [], grades, {"工商管理": courses}, {}, [], "4-1",
                entry_year=2023)


def _build_report(grade: str = "2023级") -> tuple[str, object, list]:
    from academicwarning import export

    prep = _make_prep()
    rows = check_selection_rationality(prep, "2026-09-26 00:00:00.000000")
    assert rows, "前置：构造的数据应至少触发一名学生"
    path = export.build_report_xlsx(prep, rows, grade)
    wb = load_workbook(path)
    return path, wb, rows


def test_build_report_xlsx_three_sheets():
    path, wb, _ = _build_report()

    assert os.path.isfile(path)
    assert os.path.basename(path).startswith("选课检查名单-2023级-")
    assert wb.sheetnames == ["汇总", "学生类别明细", "全类别总览"]

    ws = wb["汇总"]
    assert [c.value for c in ws[1][:4]] == ["学号", "姓名", "专业", "触发类别"]
    # 触发的 S1 在汇总首行数据区
    data_rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert any(r[0] == "S1" and r[1] == "甲" for r in data_rows)
    # 缺修必修课按类别归拢展示在"缺修课程（按类别）"列
    s1_row = next(r for r in data_rows if r[0] == "S1")
    assert "专业核心课程：" in s1_row[5]
    assert "《核心课甲》" in s1_row[5]


def test_sheet2_contains_category_detail_and_red_fill():
    _, wb, _ = _build_report()
    ws2 = wb["学生类别明细"]

    header = [c.value for c in ws2[1]]
    assert header[:5] == ["学号", "姓名", "专业", "课程类别", "应修"]
    # 表头批注（及格制口径说明）存在
    assert ws2["F1"].comment is not None

    body = [[c.value for c in r] for r in ws2.iter_rows(min_row=2)]
    # S1 专业选修：应修 4 / 已修 0 / 已选 0
    row = next(r for r in body if r[0] == "S1" and r[3] == "专业选修课程")
    assert row[4] == 4.0 and row[5] == 0.0 and row[6] == 0.0
    # 差额 > 0 的行标红填充
    row_idx = body.index(row) + 2
    assert ws2.cell(row=row_idx, column=5).fill.fgColor.rgb.endswith("FDE9E9")


def test_sheet3_overview_per_category():
    _, wb, _ = _build_report()
    ws3 = wb["全类别总览"]

    body = [[c.value for c in r] for r in ws3.iter_rows(min_row=2)]
    row = next(r for r in body if r[0] == "S1" and r[3] == "专业选修课程")
    assert row[4] == 4.0
    assert row[7] == 4.0   # 差额 = 应修 − 已修 − 已选


def test_legacy_rows_fall_back_to_message_and_recalc():
    """老批次（details 空/损坏）兜底：Sheet1 缺修列从 message 提取；
    Sheet2 缺修行用实时计算兜底（_missing_required_courses）并标红。"""
    from academicwarning import export

    prep = _make_prep()
    rows = check_selection_rationality(prep, "2026-09-26 00:00:00.000000")
    for r in rows:   # 模拟老批次：无 details 结构
        r.details = "{}"
    path = export.build_report_xlsx(prep, rows, "2023级")
    wb = load_workbook(path)

    ws = wb["汇总"]
    s1 = next(r for r in ws.iter_rows(min_row=2) if r[0].value == "S1")
    assert s1[5].value == "《核心课甲》"   # 从 message 正则提取

    ws2 = wb["学生类别明细"]
    body = [[c.value for c in r] for r in ws2.iter_rows(min_row=2)]
    miss_row = next(r for r in body if r[3] == "缺修课程（按类别）")
    assert miss_row[8] == "《核心课甲》"   # 实时计算兜底
    row_idx = body.index(miss_row) + 2
    assert ws2.cell(row=row_idx, column=5).fill.fgColor.rgb.endswith("FDE9E9")


def test_waiver_credit_overlay_and_plain_grade_filename():
    """N3 类型B 学分认可以叠层计入已修（Sheet2/3 同口径）；grade 为空 →
    文件名无年级段（选课检查名单-YYYYMMDD-HHMMSS.xlsx）。"""
    from academicwarning import export

    prep = _make_prep()
    rows = check_selection_rationality(prep, "2026-09-26 00:00:00.000000")
    prep.waiver_credit["S1"] = {"专业选修课程": 2.0}
    path = export.build_report_xlsx(prep, rows, "")
    assert os.path.basename(path).startswith("选课检查名单-")
    import re as _re
    assert not _re.search(r"-\d{4}级-", os.path.basename(path))

    wb = load_workbook(path)
    for sheet in ("学生类别明细", "全类别总览"):
        body = [[c.value for c in r]
                for r in wb[sheet].iter_rows(min_row=2)]
        row = next(r for r in body if r[0] == "S1" and r[3] == "专业选修课程")
        assert row[5] == 2.0   # 已修 = 0 + 认可 2.0；差额 = 4 − 2 = 2
        assert row[7] == 2.0


def _touch(path: str, mtime: float) -> None:
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as f:
        f.write(b"x")
    os.utime(path, (mtime, mtime))


def test_prune_reports_keeps_most_recent_five():
    from academicwarning import export

    base = 1_000_000.0
    names = [f"选课检查名单-2023级-{i:08d}.xlsx" for i in range(8)]
    for i, n in enumerate(names):
        _touch(os.path.join(export.OUT_DIR, n), base + i)
    # 干扰项：无年级段文件不应被年级清理误删
    plain = os.path.join(export.OUT_DIR, "选课检查名单-20260101-000000.xlsx")
    _touch(plain, base)

    _prune_reports("2023级")

    remaining = [os.path.basename(p)
                 for p in glob.glob(os.path.join(export.OUT_DIR, "选课检查名单-2023级-*.xlsx"))]
    # 保留最新 5 份（mtime 最大的 i=3..7）
    expected = {f"选课检查名单-2023级-{i:08d}.xlsx" for i in range(3, 8)}
    assert set(remaining) == expected
    # 无年级段文件未被清理
    assert os.path.isfile(plain)


def test_prune_reports_plain_matches_exclude_graded():
    from academicwarning import export

    base = 1_000_000.0
    plain_names = [f"选课检查名单-{i:08d}.xlsx" for i in range(7)]
    for i, n in enumerate(plain_names):
        _touch(os.path.join(export.OUT_DIR, n), base + i)
    graded = os.path.join(export.OUT_DIR, "选课检查名单-2022级-00000000.xlsx")
    _touch(graded, base)

    _prune_reports("")

    remaining = {os.path.basename(p)
                 for p in glob.glob(os.path.join(export.OUT_DIR, "选课检查名单-*.xlsx"))}
    # 无年级段保留最新 5 份；年级段文件不受影响
    expected = {f"选课检查名单-{i:08d}.xlsx" for i in range(2, 7)} | {
        "选课检查名单-2022级-00000000.xlsx"}
    assert remaining == expected
