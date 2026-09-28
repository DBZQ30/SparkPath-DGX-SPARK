"""trainingplan parse_course_notes 合成 docx 测试。

真实样本（data/SmartGuide）不随仓库分发（test_trainingplan.py 全 skip），
据此用 python-docx 内存构造覆盖分支：列定位（无编码列/无备注列/无中文
课程名称列）、空备注跳过、无编码行的 name: 兜底键、多个表格合并、
短行越界防御（note_i ≥ len(cells)）。
"""
from types import SimpleNamespace

from docx import Document

from trainingplan.parsers import _collect_row_notes, _note_columns, parse_course_notes


def _fill(table, rows):
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.cell(ri, ci).text = val


# ── _note_columns：表头列定位 ──────────────────────────────────────


def test_note_columns_locates_all_three():
    assert _note_columns(["课程编码", "中文课程名称", "学分", "备注"]) == (0, 3, 1)


def test_note_columns_missing_optional_or_required():
    # 缺中文课程名称列 → name_i=None（兜底键退化为仅编码）
    assert _note_columns(["课程编码", "备注"]) == (0, 1, None)
    # 缺备注列 / 缺编码列 → 全 None
    assert _note_columns(["课程编码", "名称"]) == (None, None, None)
    assert _note_columns(["名称", "备注"]) == (None, None, None)


# ── parse_course_notes：合成 docx 全链路 ──────────────────────────


def test_parse_course_notes_synthetic(tmp_path):
    doc = Document()
    doc.add_paragraph("课程总表")
    t1 = doc.add_table(rows=3, cols=4)
    _fill(t1, [
        ["课程编码", "中文课程名称", "学分", "备注"],
        ["082038", "管理研究方法论I", "2", "研究生进阶"],          # 编码+名称 → 两键
        ["", "前沿交叉研讨", "1", " 前沿 交叉 "],                   # 无编码 → name: 兜底
    ])
    t2 = doc.add_table(rows=2, cols=4)
    _fill(t2, [
        ["课程编码", "中文课程名称", "学分", "备注"],
        ["MAGT420208", "交叉融合课程", "2", "前沿交叉；教学改革"],  # 第二表合并入同一 notes
    ])
    # 无关表（表头缺备注列）跳过
    doc.add_table(rows=2, cols=2)
    p = tmp_path / "notes.docx"
    doc.save(p)

    notes = parse_course_notes(str(p))
    assert notes["082038"] == "研究生进阶"
    assert notes["name:管理研究方法论I"] == "研究生进阶"
    assert notes["name:前沿交叉研讨"] == "前沿交叉"   # _clean 去空白
    assert notes["MAGT420208"] == "前沿交叉；教学改革"


def test_parse_course_notes_skips_empty_note_rows(tmp_path):
    doc = Document()
    t = doc.add_table(rows=3, cols=4)
    _fill(t, [
        ["课程编码", "中文课程名称", "学分", "备注"],
        ["082038", "某课", "2", "  "],        # 备注仅空白（_clean 后为空）→ 跳过
        ["082039", "另课", "2", "双语"],      # 正常
    ])
    p = tmp_path / "notes.docx"
    doc.save(p)
    notes = parse_course_notes(str(p))
    assert notes == {"082039": "双语", "name:另课": "双语"}


# ── _collect_row_notes：短行越界防御（真实 docx 表格行等宽，构造 fake 行）──


def test_collect_row_notes_short_row_defense():
    # note_i=1：备注在第二列，短行（cells 数不足 note_i+1）跳过不抛
    rows = [
        SimpleNamespace(cells=[SimpleNamespace(text="课程编码")]),
        # 正常行：编码 + 备注（+多出的名称列，name_i=None 时不读）
        SimpleNamespace(cells=[SimpleNamespace(text="082038"),
                               SimpleNamespace(text="研究生进阶"),
                               SimpleNamespace(text="甲课")]),
        # 越界行：仅 1 个 cell → note_i=1 越界，跳过不抛
        SimpleNamespace(cells=[SimpleNamespace(text="082039")]),
        # 正常行：编码 + 备注
        SimpleNamespace(cells=[SimpleNamespace(text="082040"),
                               SimpleNamespace(text="前沿交叉")]),
    ]
    notes: dict[str, str] = {}
    _collect_row_notes(SimpleNamespace(rows=rows), 0, 1, None, notes)
    assert notes == {"082038": "研究生进阶", "082040": "前沿交叉"}
