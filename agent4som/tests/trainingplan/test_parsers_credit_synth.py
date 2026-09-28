"""trainingplan parse_credit_structure 合成 docx 测试。

真实样本（data/SmartGuide）不随仓库分发（test_trainingplan.py 全 skip），
据此用 python-docx 内存构造覆盖各分支：表头/空行/无小计行过滤、合并单元
格标签去重、带括号说明的首标签归一、连续同小计分组的公共前缀截断
（合并小计命名）、毕业要求/集中实践/课外实践顶层抽取、课程教学顶层补齐。
"""
from docx import Document

from trainingplan.parsers import parse_credit_structure

# 列：标签 ×4（合并单元格重复文本）| 小计 | 占比


def _fill(table, rows):
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            table.cell(ri, ci).text = val


def _full_rows():
    return [
        ["课程类别", "", "", "", "毕业要求学分", "占毕业总学分比例"],   # 表头行剔除
        ["毕业要求", "毕业要求", "总学分", "总学分", "148", "100%"],
        # 首标签带括号说明 → cat_norm 剥离后归"课程教学"（前缀≥2 → 二级命名）
        ["课程教学（含各类课程）", "课程教学（含各类课程）", "通识教育课程",
         "通识教育课程", "40", "27%"],
        # 连续同小计 + 同首标签 → 一组；二级标签不同 → 公共前缀截断到 1
        ["课程教学", "课程教学", "专业大类基础课程", "专业大类基础课程", "45", "30%"],
        ["课程教学", "课程教学", "专业课程", "专业课程", "45", "30%"],
        ["集中实践", "集中实践", "集中实践", "集中实践", "22", "14.9%"],
        ["课外实践", "课外实践", "课外实践", "课外实践", "8", ""],
        ["备注", "", "", "", "", ""],   # 无小计 → 剔除
        ["", "", "", "", "", ""],       # 全空行 → 剔除
    ]


def _doc_with(rows):
    doc = Document()
    t = doc.add_table(rows=len(rows), cols=6)
    _fill(t, rows)
    return doc


def test_parse_credit_structure_synth_full(tmp_path):
    """全要素：顶层三数字 + course_teaching 推导、节点结构/占比批注、
    合并小计命名（二级标签 " + " 连接）、表头与无效行过滤。"""
    path = tmp_path / "p.docx"
    _doc_with(_full_rows()).save(path)
    nodes, top = parse_credit_structure(str(path))

    assert top == {"total": "148", "practice": 22.0, "extra_practice": 8.0,
                   "course_teaching": 126.0}
    got = [(n.path, n.credit, n.category, n.credit_note, n.sort) for n in nodes]
    assert got == [
        ("毕业总学分", "148", "毕业要求", "", 0),
        ("课程教学/通识教育课程", "40", "课程教学", "占比27%", 1),
        ("课程教学/专业大类基础课程 + 专业课程", "45", "课程教学", "占比30%", 2),
        ("集中实践", "22", "集中实践", "占比14.9%", 3),
        ("课外实践", "8", "课外实践", "", 4),
        # 顶层补齐：方案表未单列"课程教学"小计 → 148 − 22 = 126
        ("课程教学", "126", "课程教学", "", 5),
    ]


def test_parse_credit_structure_synth_no_practice(tmp_path):
    """无集中实践分组 → practice 缺失 → course_teaching 不推导、不补齐；
    课外实践等其余分组正常入库。"""
    rows = [r for r in _full_rows()
            if r[0] not in ("集中实践",) and "课程教学" not in r[0]
            and not r[0].startswith("课程教学")]
    path = tmp_path / "p.docx"
    _doc_with(rows).save(path)
    nodes, top = parse_credit_structure(str(path))

    assert top == {"total": "148", "extra_practice": 8.0}
    assert [n.path for n in nodes] == ["毕业总学分", "课外实践"]
