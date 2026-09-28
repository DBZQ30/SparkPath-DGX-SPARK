"""scripts/reparse_grades.py 纯逻辑测试（此前零覆盖，离线）。

不触生产库与生产 docx：WarningDB 用 tmp_path 临时库，docx 用 python-docx
内存构造；process_file 全链路依赖真实成绩单表样，不在本地覆盖（重解析
是否成功由预览/执行的 SQL 与 diff 摘要函数保证）。
"""
import importlib.util
from pathlib import Path

import pytest
from docx import Document

from academicwarning.db import WarningDB
from academicwarning.models import Grade

_SRC = Path(__file__).resolve().parents[2] / "scripts" / "reparse_grades.py"
_spec = importlib.util.spec_from_file_location("reparse_grades_under_test", _SRC)
rp = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(rp)


class _FakeSF:
    """source_file 桩：process_file 只读 id/file_path/in_file_meta/file_name。"""

    def __init__(self, id, file_path, file_name="成绩单.docx", in_file_meta=None):
        self.id = id
        self.file_path = file_path
        self.file_name = file_name
        self.in_file_meta = in_file_meta


# ── _q：SQL 预览字面量转义 ──────────────────────────────────────────


def test_q_sql_literals():
    assert rp._q(None) == "NULL"
    assert rp._q(3) == "3"
    assert rp._q(True) == "1"
    assert rp._q(2.5) == "2.5"
    assert rp._q("张'三") == "'张''三'"   # 单引号翻倍


# ── resolve_docx_path ──────────────────────────────────────────────


def test_resolve_docx_path_absolute_and_missing(tmp_path):
    f = tmp_path / "a.docx"
    f.write_bytes(b"x")
    assert rp.resolve_docx_path(str(f)) == f
    with pytest.raises(FileNotFoundError):
        rp.resolve_docx_path(str(tmp_path / "不存在.docx"))


# ── load_students：表序映射前置 ────────────────────────────────────


def _db_with_grades(tmp_path, grades):
    db = WarningDB(str(tmp_path / "w.db"))
    db.init_schema()
    db.insert_grades(grades)
    return db


def _g(sid, name, course, fid, term="第一学年（2023-2024）"):
    return Grade(student_id=sid, student_name=name, term_label=term,
                 course_name=course, course_name_clean=course,
                 credit=2.0, grade_raw="85", source_file_id=fid)


def test_load_students_continuous_groups(tmp_path):
    db = _db_with_grades(tmp_path, [
        _g("S1", "甲", "高数", 1), _g("S1", "甲", "大物", 1),
        _g("S2", "乙", "高数", 1), _g("S2", "乙", "线代", 1),
    ])
    order, segments, total = rp.load_students(db, 1)
    assert order == [("S1", "甲"), ("S2", "乙")]
    assert segments == 2 and total == 4   # 组段 == 学生数 → 表序映射可靠


def test_load_students_interleaved_detector(tmp_path):
    """同学生不连续（S1,S2,S1）→ 段数 3 ≠ 学生 2，调用方据此跳过。"""
    db = _db_with_grades(tmp_path, [
        _g("S1", "甲", "高数", 1), _g("S2", "乙", "大物", 1), _g("S1", "甲", "线代", 1),
    ])
    order, segments, total = rp.load_students(db, 1)
    assert order == [("S1", "甲"), ("S2", "乙")]
    assert segments == 3 and total == 3


# ── student_diff ───────────────────────────────────────────────────


def test_student_diff_semantics():
    old = [_g("S1", "甲", "高数", 1, term="T1"), _g("S1", "甲", "旧课", 1, term="T1")]
    new = [_g("S1", "甲", "高数", 1, term="T2"), _g("S1", "甲", "新课", 1, term="T2")]
    d = rp.student_diff(old, new)
    assert d == {"old_n": 2, "new_n": 2, "changed": 1,   # 高数 term T1→T2
                 "recovered": ["新课"], "lost": ["旧课"]}


# ── _map_tables_to_students：两个跳过分支 ──────────────────────────


def test_map_tables_skip_on_table_count_mismatch(tmp_path):
    db = _db_with_grades(tmp_path, [_g("S1", "甲", "高数", 7)])
    doc = Document()  # 0 张表
    p = tmp_path / "s.docx"
    doc.save(p)
    import io
    out = io.StringIO()
    assert rp._map_tables_to_students(db, _FakeSF(7, str(p)), out) is None
    assert "跳过：学生数 1 ≠ 成绩单表数 0" in out.getvalue()


def test_map_tables_skip_on_discontinuous_groups(tmp_path):
    """学生组不连续 → 段数校验拒绝（表数读取之前）。"""
    db = _db_with_grades(tmp_path, [
        _g("S1", "甲", "高数", 7), _g("S2", "乙", "大物", 7), _g("S1", "甲", "线代", 7),
    ])
    p = tmp_path / "any.docx"
    Document().save(p)   # 内容无关——段数校验在读表之前
    import io
    out = io.StringIO()
    assert rp._map_tables_to_students(db, _FakeSF(7, str(p)), out) is None
    assert "跳过：现库学生组不连续（组段 3 ≠ 学生 2）" in out.getvalue()


def test_map_tables_success(tmp_path):
    db = _db_with_grades(tmp_path, [_g("S1", "甲", "高数", 7), _g("S1", "甲", "大物", 7)])
    doc = Document()
    doc.add_table(rows=2, cols=3)
    p = tmp_path / "s.docx"
    doc.save(p)
    import io
    out = io.StringIO()
    mapped = rp._map_tables_to_students(db, _FakeSF(7, str(p)), out)
    assert mapped is not None
    orders, tables, old_total = mapped
    assert orders == [("S1", "甲")] and len(tables) == 1 and old_total == 2
    assert "按表序映射学生 ✓" in out.getvalue()


# ── _group_old_by_student ──────────────────────────────────────────


def test_group_old_by_student_segments_in_order(tmp_path):
    db = _db_with_grades(tmp_path, [
        _g("S1", "甲", "高数", 7), _g("S1", "甲", "大物", 7),
        _g("S2", "乙", "线代", 7),
    ])
    orders = [("S1", "甲"), ("S2", "乙")]
    grouped = rp._group_old_by_student(db, _FakeSF(7, "x"), orders)
    assert [g.course_name for g in grouped[0]] == ["高数", "大物"]
    assert [g.course_name for g in grouped[1]] == ["线代"]


# ── _build_sql_lines / process_file 预览出口 ───────────────────────


def test_build_sql_lines_preview(tmp_path):
    sf = _FakeSF(9, "x", in_file_meta={"major": "工商管理"})
    rows = [_g("S1", "甲", "高数", 9), _g("S1", "甲", "张'三", 9)]
    lines, row_cnt = rp._build_sql_lines(sf, [rows])
    assert row_cnt == 2
    assert lines[0] == "-- source_file #9 工商管理 成绩单.docx"
    assert lines[1] == "DELETE FROM grade WHERE source_file_id = 9;"
    assert "INSERT INTO grade (student_id," in lines[2]
    assert "'张''三'" in lines[3]   # 预览输出同样转义单引号
